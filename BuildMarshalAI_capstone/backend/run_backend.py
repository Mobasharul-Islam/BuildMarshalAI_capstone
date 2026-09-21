from __future__ import annotations
# Auto-generated backend runner for BuildMarshalAI
import os, sys
from pathlib import Path

# Force UTF-8 stdout and stderr encoding on Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

backend_dir = Path(__file__).resolve().parent
repo_root = backend_dir.parent
sys.path.insert(0, str(backend_dir))
sys.path.insert(0, str(repo_root))
os.chdir(str(backend_dir))


# ======================================================================
# CELL 0
# ======================================================================

# STEP 0 - Runtime and Hugging Face configuration
import os, re, subprocess, sys
from pathlib import Path

# Configure UTF-8 streams on Windows to prevent UnicodeEncodeError with emoji outputs
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

if os.name != "nt":
    os.environ["PYTORCH_CUDA_ALLOC_CONF"] = "expandable_segments:True"
os.environ["HF_HUB_DISABLE_XET"] = "1"
os.environ["HF_HUB_DOWNLOAD_TIMEOUT"] = "120"
os.environ["HF_HUB_ETAG_TIMEOUT"] = "60"

# Prevent torchao / torch.int1 mismatch on PyTorch < 2.6
sys.modules.setdefault("torchao", None)

IS_KAGGLE = Path("/kaggle/working").exists()
_cwd = Path.cwd()
if (_cwd / "BuildMarshalAI_capstone").exists():
    _local_repo_root = _cwd / "BuildMarshalAI_capstone"
elif _cwd.name == "backend" and (_cwd.parent / "frontend").exists():
    _local_repo_root = _cwd.parent
else:
    _local_repo_root = _cwd
RUNTIME_ROOT = Path("/kaggle/working") if IS_KAGGLE else _local_repo_root / ".buildmarshal_runtime"
RUNTIME_ROOT.mkdir(parents=True, exist_ok=True)
os.environ.setdefault("HF_HOME", str(RUNTIME_ROOT / "hf_cache"))
os.environ.setdefault("PIP_CACHE_DIR", str(RUNTIME_ROOT / "pip_cache"))
# Kaggle can discard large weights after loading to conserve session disk.
# A local backend keeps the verified files so later starts do not redownload them.
os.environ.setdefault("BUILDMARSHAL_KEEP_MODEL_FILES", "0" if IS_KAGGLE else "1")

# Load secrets without printing their values. Local Windows runs use environment variables.
if IS_KAGGLE:
    try:
        from kaggle_secrets import UserSecretsClient
        _secrets = UserSecretsClient()
        for _name in ("HF_TOKEN", "NGROK_AUTH_TOKEN", "CLIPROXY_BASE_URL", "CLIPROXY_API_KEY", "CLIPROXY_MODEL"):
            try:
                _value = _secrets.get_secret(_name)
                if _value:
                    os.environ[_name] = _value
            except Exception:
                pass
        print("Secrets loaded from Kaggle Secrets where available")
    except Exception as _secret_error:
        print(f"Kaggle Secrets unavailable: {_secret_error}")
else:
    # Reuse the API key and the port already configured for the local
    # CLIProxyAPI service. Values are loaded into memory only and are never
    # printed or copied into this notebook, so a fresh VS Code kernel needs
    # no manual environment setup.
    #
    # The port is read rather than assumed: Windows hands blocks of low TCP
    # ports to Hyper-V at every boot and then refuses to bind them, so the
    # proxy does not always end up on 8317, and a notebook started by hand
    # does not inherit the launcher's environment variables. Reading the
    # same file the proxy was started with is what keeps the two agreeing.
    _from = Path(globals()["__file__"]).resolve().parent if "__file__" in globals() else Path.cwd()
    _proxy_candidates = [_parent / "cliproxyapi" / "config.yaml"
                         for _parent in (_from, *_from.parents[:4])]
    _proxy_candidates.append(Path.home() / ".cli-proxy-api" / "config.yaml")
    _proxy_port = 8317
    for _proxy_config in _proxy_candidates:
        if not _proxy_config.exists():
            continue
        _config_text = _proxy_config.read_text(encoding="utf-8")
        _key_match = re.search(r"(?m)^api-keys:\s*\r?\n\s*-\s*[\"']?([^\"'\r\n#]+)", _config_text)
        if _key_match and not os.environ.get("CLIPROXY_API_KEY"):
            os.environ["CLIPROXY_API_KEY"] = _key_match.group(1).strip()
        _port_match = re.search(r"(?m)^\s*port:\s*(\d+)", _config_text)
        if _port_match:
            _proxy_port = int(_port_match.group(1))
        break
    # setdefault is not enough: PREPARE-TEAMMATE-PC.ps1 persists
    # CLIPROXY_BASE_URL as a user environment variable, so a stale port
    # survives in every shell and kernel and quietly outranks the config the
    # proxy was actually started from. A loopback address that disagrees
    # with that config is wrong by construction; a remote one is deliberate.
    _current = os.environ.get("CLIPROXY_BASE_URL", "").strip()
    _wanted = f"http://127.0.0.1:{_proxy_port}/v1"
    _is_loopback = re.match(r"https?://(127\.0\.0\.1|localhost|\[::1\])(:|/|$)",
                            _current, re.I)
    if not _current or (_is_loopback and _current.rstrip("/") != _wanted.rstrip("/")):
        if _current:
            print(f"CLIPROXY_BASE_URL was {_current}, but the proxy is configured "
                  f"for port {_proxy_port}; using {_wanted}.")
        os.environ["CLIPROXY_BASE_URL"] = _wanted
    os.environ.setdefault("CLIPROXY_MODEL", "gemini-3.7-flash-high")

# Keep Kaggle's CUDA wheel repair, but never replace the user's local Windows torch install.
if IS_KAGGLE:
    subprocess.run([
        sys.executable, "-m", "pip", "install", "-q", "--no-cache-dir", "--upgrade",
        "torch", "torchvision", "torchaudio", "--index-url",
        "https://download.pytorch.org/whl/cu121",
    ], check=True)

import torch
print("PyTorch:", torch.__version__)
print("CUDA available:", torch.cuda.is_available())
if not torch.cuda.is_available():
    raise RuntimeError("A CUDA GPU is required for the ColPali retriever.")
for i in range(torch.cuda.device_count()):
    props = torch.cuda.get_device_properties(i)
    print(f"  [cuda:{i}] {props.name} - {props.total_memory/1e9:.1f} GB")


# ======================================================================
# CELL 3
# ======================================================================

import os, re, json, uuid, time, shutil, base64, asyncio, logging, sys
from pathlib import Path
from typing import Optional, List, Dict, Any
from io import BytesIO
from datetime import datetime, timezone

try:
    import asyncio
    asyncio.get_running_loop()
    import nest_asyncio
    nest_asyncio.apply()
except (RuntimeError, NameError, ImportError):
    pass

import torch
import pandas as pd
from PIL import Image
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from fastapi import Depends, FastAPI, UploadFile, File, Form, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
import uvicorn
from pyngrok import ngrok

logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s")
logger = logging.getLogger("BuildMarshalAI")

# BASE_DIR holds only account-independent state: the shared model cache and the
# accounts/, accounts.json, users.json, and sessions.json registry files.  Every
# document, page image, vector index, and JSON store belongs to exactly one
# account and lives under BASE_DIR/accounts/<account_id>/, reached through the
# AccountWorkspace resolved from the caller's bearer token.
BASE_DIR = Path(os.environ.get("BUILDMARSHAL_DATA_DIR", str(RUNTIME_ROOT / "buildmarshal")))
ACCOUNTS_ROOT = BASE_DIR / "accounts"
MODEL_ROOT = BASE_DIR / "hf_models"
for directory in (BASE_DIR, ACCOUNTS_ROOT, MODEL_ROOT):
    directory.mkdir(parents=True, exist_ok=True)

# Upload ceiling enforced server-side; the frontend applies the same limit.
MAX_UPLOAD_BYTES = int(os.environ.get("BUILDMARSHAL_MAX_UPLOAD_BYTES", str(50 * 1024 * 1024)))

# Browsers must be told explicitly which origins may call the API.  A wildcard
# origin is invalid whenever credentials are allowed, and this backend now
# carries bearer tokens, so the allow-list is configurable and closed by default.
ALLOWED_ORIGINS = [
    origin.strip().rstrip("/")
    for origin in os.environ.get(
        "BUILDMARSHAL_ALLOWED_ORIGINS",
        "http://localhost:5500,http://127.0.0.1:5500,http://localhost:3000,http://127.0.0.1:3000",
    ).split(",")
    if origin.strip()
]

NGROK_AUTH_TOKEN = os.environ.get("NGROK_AUTH_TOKEN", "").strip()

NUM_GPUS = torch.cuda.device_count()
COLPALI_DEVICE_IDX = 0
COLPALI_DEVICE = "cuda:0"
print(f"Config ready | GPUs detected: {NUM_GPUS}")
print(f"ColPali retrieval -> {COLPALI_DEVICE}; generation -> CLIProxyAPI at "
      f"{os.environ.get('CLIPROXY_BASE_URL', 'http://127.0.0.1:8317/v1')}")


# ======================================================================
# CELL 4
# ======================================================================

# =======================================================================
# CELL 5 — ColPali  +  ChromaDB vector store
# =======================================================================

import gc, torch, numpy as np
import chromadb
from typing import Optional, List, Dict
from PIL import Image

# ── PEFT / transformers KeyError: 'llava' patch ──
# PaliGemma uses a llava-style architecture. When transformers' PEFT integration
# tries to look it up in _MOE_TARGET_MODULE_MAPPING it fails with KeyError: 'llava'.
# These two patches make that lookup a safe no-op.
import transformers.integrations.peft as _peft_integ

if hasattr(_peft_integ, "_convert_peft_config_moe"):
    _orig_moe = _peft_integ._convert_peft_config_moe
    def _safe_moe(p, m):
        try: return _orig_moe(p, m)
        except KeyError: return p
    _peft_integ._convert_peft_config_moe = _safe_moe

if hasattr(_peft_integ, "convert_peft_config_for_transformers"):
    _orig_cvt = _peft_integ.convert_peft_config_for_transformers
    def _safe_cvt(p, model=None, conversions=None):
        try: return _orig_cvt(p, model=model, conversions=conversions)
        except KeyError: return p
    _peft_integ.convert_peft_config_for_transformers = _safe_cvt

print("✅ PEFT patch applied (or not needed on this transformers version)")

# ── Free VRAM before loading ──
for _v in ['COLPALI_MODEL', 'COLPALI_PROCESSOR']:
    if _v in dir() and eval(_v) is not None:
        exec(f"del {_v}")
gc.collect()
torch.cuda.empty_cache()
print(f"Free VRAM on {COLPALI_DEVICE} before load: {torch.cuda.mem_get_info(COLPALI_DEVICE_IDX)[0]/1e9:.2f} GB")

# ── PaliGemma 'language_model' compatibility patch ──
# colpali_engine==0.3.1 does `model.language_model._tied_weights_keys`,
# but newer transformers wraps PaliGemma's language model under
# `.model.language_model` instead of exposing `.language_model` directly.
import sys
sys.modules.setdefault('torchao', None)
from transformers import PaliGemmaForConditionalGeneration

def _colpali_compat_language_model(self):
    direct = self._modules.get("language_model")
    if direct is not None:
        return direct
    return self.model.language_model

PaliGemmaForConditionalGeneration.language_model = property(_colpali_compat_language_model)
print("✅ PaliGemma compatibility patch applied")

# ── Load ColPali using colpali_engine's own classes ──
from colpali_engine.models import ColPali, ColPaliProcessor

ADAPTER_ID = "vidore/colpali-v1.2"

# Kaggle's current huggingface_hub route can receive invalid signed URLs from
# us.gcp.cdn.hf.co. Download the exact inference files through Hugging Face's
# public resolve endpoint instead; this redirects to the working CAS bridge.
# Model revisions and SHA-256 values are pinned for reproducibility and safety.
import hashlib
import time as _download_time
import requests

# MODEL_ROOT is defined with the other shared paths in the configuration cell.
ADAPTER_DIR = MODEL_ROOT / "colpali-v1.2"
BASE_MODEL_ID = "vidore/colpaligemma-3b-pt-448-base"
BASE_MODEL_DIR = MODEL_ROOT / "colpaligemma-3b-pt-448-base"
MODEL_ROOT.mkdir(parents=True, exist_ok=True)

ADAPTER_REVISION = "6b89bc63c16809af4d111bfe412e2ac6bc3c9451"
BASE_REVISION = "30ab955d073de4a91dc5a288e8c97226647e3e5a"
SIGNED_URL_OVERRIDES = {}  # populated only by short-lived Kaggle validation runs


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(16 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _download_resolve_file(repo_id, revision, filename, destination, expected_size, expected_sha=None):
    destination.parent.mkdir(parents=True, exist_ok=True)
    locally_rewritten_adapter = (
        filename == "adapter_config.json"
        and destination.exists()
        and os.path.normcase(os.path.abspath(
            json.loads(destination.read_text(encoding="utf-8")).get("base_model_name_or_path", "")
        )) == os.path.normcase(os.path.abspath(str(BASE_MODEL_DIR)))
    )
    if locally_rewritten_adapter or (destination.exists() and destination.stat().st_size == expected_size):
        if expected_sha is None or _sha256(destination) == expected_sha:
            logger.info(f"Using verified cached file: {filename}")
            return
        destination.unlink()

    partial = destination.with_name(destination.name + ".part")
    for attempt in range(1, 11):
        resume_at = partial.stat().st_size if partial.exists() else 0
        headers = {"Range": f"bytes={resume_at}-"} if resume_at else {}
        if os.environ.get("HF_TOKEN"):
            headers["Authorization"] = f"Bearer {os.environ['HF_TOKEN']}"
        url = SIGNED_URL_OVERRIDES.get((repo_id, filename)) or (
            f"https://huggingface.co/{repo_id}/resolve/{revision}/{filename}"
        )
        try:
            logger.info(
                f"Direct download {repo_id}/{filename} attempt {attempt}/10 "
                f"(resume={resume_at/1e6:.1f} MB)"
            )
            with requests.get(
                url,
                headers=headers,
                stream=True,
                allow_redirects=True,
                timeout=(30, 120),
            ) as response:
                response.raise_for_status()
                if resume_at and response.status_code != 206:
                    partial.unlink(missing_ok=True)
                    resume_at = 0
                mode = "ab" if resume_at and response.status_code == 206 else "wb"
                with partial.open(mode) as handle:
                    for chunk in response.iter_content(chunk_size=8 * 1024 * 1024):
                        if chunk:
                            handle.write(chunk)
            if partial.stat().st_size != expected_size:
                raise RuntimeError(
                    f"Size mismatch for {filename}: {partial.stat().st_size} != {expected_size}"
                )
            partial.replace(destination)
            if expected_sha and _sha256(destination) != expected_sha:
                destination.unlink(missing_ok=True)
                raise RuntimeError(f"SHA-256 mismatch for {filename}")
            logger.info(f"Verified download: {filename} ({expected_size/1e9:.2f} GB)")
            return
        except Exception as error:
            logger.warning(f"Direct download failed for {filename}: {error}")
            if attempt == 10:
                raise
            _download_time.sleep(min(2 * attempt, 15))


def _download_repo_files(repo_id, revision, destination_dir, files):
    for filename, expected_size, expected_sha in files:
        _download_resolve_file(
            repo_id,
            revision,
            filename,
            destination_dir / filename,
            expected_size,
            expected_sha,
        )


_download_repo_files(
    ADAPTER_ID,
    ADAPTER_REVISION,
    ADAPTER_DIR,
    [
        ("adapter_config.json", 750, None),
        ("adapter_model.safetensors", 78625112, "caed65068cae6d50e572d984914324a7d8a9360cdd7f4263ea82f1792614391f"),
        ("preprocessor_config.json", 700, None),
        ("special_tokens_map.json", 733, None),
        ("tokenizer.json", 17763459, "ffd310e50986db7a039948ab83441d612689e7f989198e31b5c8984ca458adf6"),
        ("tokenizer_config.json", 242696, None),
    ],
)

_download_repo_files(
    BASE_MODEL_ID,
    BASE_REVISION,
    BASE_MODEL_DIR,
    [
        ("config.json", 1015, None),
        ("model.safetensors.index.json", 66301, None),
        ("model-00001-of-00002.safetensors", 4986817288, "c128f5670d7a66942a194be6e2d324dc329c0de19e99c6f047513878e14f988e"),
        ("model-00002-of-00002.safetensors", 862495528, "8352c38e4d1785c4a35547d13f4d8d5562faab6fe8e9a30b1f5d8039d355a409"),
    ],
)

# Make PEFT resolve the base model locally. This prevents from_pretrained()
# from issuing another implicit Hugging Face download for the two large shards.
_adapter_config_path = ADAPTER_DIR / "adapter_config.json"
_adapter_config = json.loads(_adapter_config_path.read_text(encoding="utf-8"))
_adapter_config["base_model_name_or_path"] = str(BASE_MODEL_DIR)
_adapter_config_path.write_text(
    json.dumps(_adapter_config, indent=2), encoding="utf-8"
)
MODEL_SNAPSHOT = str(ADAPTER_DIR)
logger.info(f"ColPali inference files ready: {MODEL_SNAPSHOT}")

logger.info(f"Loading ColPali v1.2 onto {COLPALI_DEVICE} from local snapshot (fp16)...")
COLPALI_MODEL = ColPali.from_pretrained(
    MODEL_SNAPSHOT,
    local_files_only=True,
    torch_dtype=torch.float16,
    device_map={"": COLPALI_DEVICE_IDX},
    attn_implementation="sdpa",
    low_cpu_mem_usage=True,
).eval()

COLPALI_PROCESSOR = ColPaliProcessor.from_pretrained(
    MODEL_SNAPSHOT, local_files_only=True
)

print(f"ColPali ready on {COLPALI_DEVICE}. Free VRAM: {torch.cuda.mem_get_info(COLPALI_DEVICE_IDX)[0]/1e9:.2f} GB")
if os.environ.get("BUILDMARSHAL_KEEP_MODEL_FILES", "0") != "1":
    shutil.rmtree(ADAPTER_DIR, ignore_errors=True)
    shutil.rmtree(BASE_MODEL_DIR, ignore_errors=True)
    logger.info("Removed loaded ColPali weight files to free Kaggle disk space")

# =======================================================================
# ChromaDB - one persistent vector store per account
# =======================================================================
#
# The ColPali model is stateless and shared by every request.  The index is
# not: each account owns a ChromaDB directory under its workspace, so a query
# physically cannot reach another account's vectors even if a metadata filter
# were ever omitted.  ``build_account_collection`` is handed to the account
# registry, which calls it lazily the first time a workspace needs its index.
# =======================================================================

_COLLECTION_NAME = "document_pages"


# ChromaDB >=0.6.x requires an explicit EmbeddingFunction object; passing
# None raises a type error.  We supply a no-op stub because ColPali
# produces all embeddings externally - ChromaDB never calls this function.
class _NoOpEF(chromadb.EmbeddingFunction):
    """Stub: ColPali supplies all vectors; ChromaDB never calls this."""
    def __call__(self, input):
        raise RuntimeError("_NoOpEF should never be called directly")

_NO_OP_EF = _NoOpEF()


def build_account_collection(chroma_dir):
    """Open (or create) one account's ``document_pages`` collection.

    Returns the client alongside the collection: the workspace has to close the
    client before its directory can be deleted, because ChromaDB holds the
    SQLite file open for the client's lifetime.
    """
    client = chromadb.PersistentClient(path=str(chroma_dir))
    try:
        collection = client.get_collection(name=_COLLECTION_NAME, embedding_function=_NO_OP_EF)
    except Exception:
        collection = client.create_collection(
            name=_COLLECTION_NAME,
            embedding_function=_NO_OP_EF,
            metadata={"hnsw:space": "cosine"},
        )
    return client, collection


@torch.no_grad()
def embed_image(image_path: str, workspace=None) -> Optional[np.ndarray]:
    try:
        # ``workspace`` is accepted for signature parity with the hybrid
        # retriever that replaces this function; pooled embedding is uncached.
        img = Image.open(image_path).convert("RGB")

        # process_images() builds the correct token structure for page screenshots
        batch = COLPALI_PROCESSOR.process_images([img]).to(COLPALI_MODEL.device)

        # ColPali.forward() returns a raw Tensor (batch, seq_len, dim) - NOT a ModelOutput
        out = COLPALI_MODEL(**batch)           # shape: (1, seq_len, embedding_dim)
        vec = out.mean(dim=1).squeeze(0)       # mean-pool -> (embedding_dim,)

        return vec.float().cpu().numpy()
    except Exception as e:
        logger.error(f"embed_image failed for {image_path}: {e}")
        return None


@torch.no_grad()
def embed_query(query: str) -> Optional[np.ndarray]:
    try:
        # process_queries() builds the correct token structure for text queries
        batch = COLPALI_PROCESSOR.process_queries([query]).to(COLPALI_MODEL.device)

        # Same - ColPali returns a raw Tensor
        out = COLPALI_MODEL(**batch)           # shape: (1, seq_len, embedding_dim)
        vec = out.mean(dim=1).squeeze(0)       # mean-pool -> (embedding_dim,)

        return vec.float().cpu().numpy()
    except Exception as e:
        logger.error(f"embed_query failed: {e}")
        return None


def get_all_pages(workspace, limit: int = 5) -> List[Dict]:
    """Fallback context for an account whose index is empty or unqueryable."""
    pages = []
    for doc_id, doc in workspace.load_metadata().get("documents", {}).items():
        for page in doc.get("pages", []):
            pages.append({
                "doc_name":     doc.get("name", doc_id),
                "doc_id":       doc_id,
                "page":         page.get("page_num", 1),
                "image_path":   workspace.resolve_page_path(page.get("image_path", "")),
                "text_content": page.get("text_content", ""),
                "score":        0.0,
            })
            if len(pages) >= limit:
                return pages
    return pages


def retrieve_context(query: str, top_k: int = 5, project_id: str | None = None,
                     workspace=None) -> List[Dict]:
    """Pooled-vector retrieval within one account's index.

    Replaced at startup by the hybrid ColPali retriever; kept as the fallback
    used when that integration is unavailable.
    """
    if workspace is None:
        raise RuntimeError("retrieve_context requires the caller's account workspace")
    collection = workspace.collection
    if collection.count() == 0:
        logger.warning("Vector index empty for this account - falling back to all pages")
        return get_all_pages(workspace, top_k)

    q_vec = embed_query(query)
    if q_vec is None:
        logger.error("Query embedding failed - falling back to all pages")
        return get_all_pages(workspace, top_k)

    kwargs = {
        "query_embeddings": [q_vec.tolist()],
        "n_results": min(top_k, collection.count()),
        "include": ["metadatas", "distances"],
    }
    if project_id:
        kwargs["where"] = {"project_id": project_id}
    results = collection.query(**kwargs)

    pages = []
    for meta, dist in zip(results["metadatas"][0], results["distances"][0]):
        pages.append({
            "doc_name":     meta.get("doc_name", ""),
            "doc_id":       meta.get("doc_id", ""),
            "page":         int(meta.get("page_num", 1)),
            "image_path":   workspace.resolve_page_path(meta.get("image_path", "")),
            "text_content": meta.get("text_content", ""),
            "score":        round(1.0 - dist, 4),
        })

    if pages:
        logger.info(f"Retrieved {len(pages)} pages (top similarity: {pages[0]['score']})")
    return pages


print("ColPali ready; per-account ChromaDB vector stores enabled")


# ======================================================================
# CELL 5
# ======================================================================

# =======================================================================
# CELL 6 — Document Ingestion  ->  ColPali embedding  ->  ChromaDB
# =======================================================================
#
# Flow for each file:
#   convert to page screenshots
#   for each page:
#       embed_image(screenshot)  -> float32 numpy vector
#       workspace.collection.add(id, embedding, metadata, document)
#   persist the account's metadata.json (for /api/documents and deletion)
# =======================================================================

def process_pdf(file_path: Path, doc_id: str, pages_dir: Path) -> List[Dict]:
    """Render PDF pages with PyMuPDF (no external Poppler installation)."""
    import fitz

    logger.info(f"Processing PDF: {file_path.name}")
    page_infos = []
    with fitz.open(file_path) as document:
        scale = 150 / 72
        matrix = fitz.Matrix(scale, scale)
        for index, pdf_page in enumerate(document):
            pixmap = pdf_page.get_pixmap(matrix=matrix, alpha=False)
            page_img = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            page_img.thumbnail((1024, 1024))
            page_path = pages_dir / f"{doc_id}_page_{index + 1}.png"
            page_img.save(page_path, "PNG")
            page_infos.append({
                "page_num": index + 1,
                "image_path": str(page_path),
                "width": page_img.width,
                "height": page_img.height,
                "text_content": pdf_page.get_text("text")[:12000],
            })
    logger.info(f"PDF: {len(page_infos)} pages extracted")
    return page_infos


def _df_to_image(df, title, out_path):
    fig, ax = plt.subplots(figsize=(min(20, len(df.columns)*1.5+1),
                                    min(20, len(df)*0.4+1)))
    ax.axis('off')
    data = df.head(50).fillna('').astype(str)
    tbl  = ax.table(cellText=data.values, colLabels=data.columns,
                    cellLoc='left', loc='center')
    tbl.auto_set_font_size(False); tbl.set_fontsize(8); tbl.scale(1, 1.3)
    for (r, c), cell in tbl.get_celld().items():
        if r == 0:
            cell.set_facecolor('#4472C4')
            cell.set_text_props(color='white', fontweight='bold')
        else:
            cell.set_facecolor('#f0f0f0' if r % 2 == 0 else 'white')
    plt.title(title, fontsize=10, fontweight='bold')
    plt.tight_layout()
    plt.savefig(str(out_path), dpi=120, bbox_inches='tight', facecolor='white')
    plt.close(fig)


def process_excel(file_path: Path, doc_id: str, pages_dir: Path) -> List[Dict]:
    page_infos = []
    dfs = ({"Sheet1": pd.read_csv(str(file_path))}
           if file_path.suffix == '.csv'
           else pd.read_excel(str(file_path), sheet_name=None))
    for sheet_name, df in dfs.items():
        page_path = pages_dir / f"{doc_id}_sheet_{sheet_name}.png"
        try:
            _df_to_image(df, f"{file_path.stem} — {sheet_name}", page_path)
            page_infos.append({
                "page_num":    len(page_infos) + 1,
                "image_path":  str(page_path),
                "sheet_name":  sheet_name,
                "text_content": df.head(100).to_string(),
            })
        except Exception as e:
            logger.warning(f"Sheet render failed: {e}")
            page_infos.append({
                "page_num":    len(page_infos) + 1,
                "image_path":  None,
                "sheet_name":  sheet_name,
                "text_content": df.head(100).to_string(),
            })
    return page_infos


def process_image(file_path: Path, doc_id: str, pages_dir: Path) -> List[Dict]:
    img = Image.open(str(file_path)).convert("RGB")
    img.thumbnail((1024, 1024))
    page_path = pages_dir / f"{doc_id}_img.png"
    img.save(str(page_path), "PNG")
    return [{
        "page_num": 1, "image_path": str(page_path),
        "width": img.width, "height": img.height, "text_content": "",
    }]


def process_docx(file_path: Path, doc_id: str, pages_dir: Path) -> List[Dict]:
    from docx import Document as DocxDocument
    doc       = DocxDocument(str(file_path))
    full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    page_infos = []
    chunks = [full_text[i:i+2000] for i in range(0, max(len(full_text), 1), 2000)]
    for idx, chunk in enumerate(chunks):
        page_path = pages_dir / f"{doc_id}_page_{idx+1}.png"
        try:
            fig, ax = plt.subplots(figsize=(10, 14)); ax.axis('off')
            ax.text(0.05, 0.95, chunk, transform=ax.transAxes, fontsize=9,
                    verticalalignment='top', fontfamily='monospace', wrap=True)
            plt.title(f"{file_path.stem} — Page {idx+1}")
            plt.savefig(str(page_path), dpi=120, bbox_inches='tight', facecolor='white')
            plt.close(fig)
            page_infos.append({
                "page_num":    idx + 1,
                "image_path":  str(page_path),
                "text_content": chunk,
            })
        except Exception as e:
            page_infos.append({
                "page_num":    idx + 1,
                "image_path":  None,
                "text_content": chunk,
            })
    return page_infos
def process_audio(file_path: Path, doc_id: str, pages_dir: Path) -> List[Dict]:
    """
    Audio files can't be visually embedded; store as a single 'page'.
    The pages endpoint will serve the raw audio file for playback.
    """
    # Copy audio to pages dir for serving
    page_path = pages_dir / f"{doc_id}_audio{file_path.suffix}"
    shutil.copy2(str(file_path), str(page_path))
    
    return [{
        "page_num": 1,
        "image_path": str(page_path),  # repurpose image_path for serving
        "text_content": f"Audio file: {file_path.name}",
    }]


def process_text(file_path: Path, doc_id: str, pages_dir: Path) -> List[Dict]:
    """Process plain text files (txt, json, xml, code files etc.)."""
    text = file_path.read_text(encoding="utf-8", errors="replace")
    page_infos = []
    chunks = [text[i:i+3000] for i in range(0, max(len(text), 1), 3000)]
    
    for idx, chunk in enumerate(chunks):
        page_path = pages_dir / f"{doc_id}_page_{idx+1}.png"
        try:
            fig, ax = plt.subplots(figsize=(10, 14)); ax.axis('off')
            ax.text(0.05, 0.95, chunk, transform=ax.transAxes, fontsize=8,
                    verticalalignment='top', fontfamily='monospace', wrap=True)
            plt.title(f"{file_path.stem} — Page {idx+1}")
            plt.savefig(str(page_path), dpi=120, bbox_inches='tight', facecolor='white')
            plt.close(fig)
            page_infos.append({
                "page_num": idx + 1,
                "image_path": str(page_path),
                "text_content": chunk,
            })
        except Exception:
            page_infos.append({
                "page_num": idx + 1,
                "image_path": None,
                "text_content": chunk,
            })
    return page_infos


def process_generic(file_path: Path, doc_id: str, pages_dir: Path) -> List[Dict]:
    """Fallback — store file metadata only."""
    page_path = pages_dir / f"{doc_id}_file{file_path.suffix}"
    shutil.copy2(str(file_path), str(page_path))
    return [{
        "page_num": 1,
        "image_path": str(page_path),
        "text_content": f"File: {file_path.name} ({file_path.stat().st_size} bytes)",
    }]


def _file_digest(path: Path, chunk: int = 1024 * 1024) -> str:
    """A content hash of an upload, read in chunks so a large file never lands
    in memory.  Matches document_storage.digest_of."""
    import hashlib

    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        while block := handle.read(chunk):
            digest.update(block)
    return digest.hexdigest()


def ingest_document(file_path: Path, doc_id: str, workspace, display_name: str | None = None) -> Dict:
    """
    Convert document to page screenshots, embed each with ColPali,
    and store the resulting vector + metadata in ChromaDB.

    ChromaDB entry per page
      id        : "<doc_id>_p<page_num>"  unique string key
      embedding : ColPali mean-pooled float32 vector (provided as list)
      metadata  : doc_id, doc_name, page_num, image_path, text_content
      document  : text_content (ChromaDB full-text field)
    """
    ext  = file_path.suffix.lower()
    # The file on disk is named after its id, so the caller supplies the name
    # the user actually uploaded.  It reaches the document list, the citations
    # in chat answers, and the lexical half of retrieval scoring.
    display_name = Path(display_name).name if display_name else file_path.name
    pages_dir = workspace.pages_dir
    collection = workspace.collection
    meta = workspace.load_metadata()

    if ext == '.pdf':
        pages = process_pdf(file_path, doc_id, pages_dir)
    elif ext in ('.xlsx', '.xls', '.csv'):
        pages = process_excel(file_path, doc_id, pages_dir)
    elif ext in ('.png', '.jpg', '.jpeg', '.gif', '.webp', '.bmp', '.tiff'):
        pages = process_image(file_path, doc_id, pages_dir)
    elif ext in ('.doc', '.docx'):
        pages = process_docx(file_path, doc_id, pages_dir)
    elif ext in ('.mp3', '.wav', '.ogg', '.m4a', '.flac', '.aac', '.wma', '.opus'):
        pages = process_audio(file_path, doc_id, pages_dir)
    elif ext in ('.txt', '.csv', '.json', '.xml', '.html', '.css', '.js', '.py', 
                 '.java', '.c', '.cpp', '.md', '.rtf'):
        pages = process_text(file_path, doc_id, pages_dir)
    else:
        pages = process_generic(file_path, doc_id, pages_dir)  # fallback instead of error

    logger.info(f"Embedding {len(pages)} pages and storing in ChromaDB...")
    embedded = 0

    for page in pages:
        img_path = page.get("image_path")
        if not img_path or not os.path.exists(img_path):
            logger.warning(f"Skipping page {page['page_num']} — no image file")
            continue

        # 1. Embed the page screenshot with ColPali
        embedding = embed_image(img_path, workspace=workspace)
        if embedding is None:
            logger.warning(f"Embedding returned None for page {page['page_num']}")
            continue

        # 2. Build ChromaDB entry
        chroma_id  = f"{doc_id}_p{page['page_num']}"
        # ChromaDB metadata values must be str / int / float / bool.
        # Truncate text_content to 1000 chars to stay within metadata limits.
        text_trunc = page.get("text_content", "")[:1000]

        page_metadata = {
            "doc_id":       doc_id,
            "doc_name":     display_name,
            "page_num":     int(page["page_num"]),
            "image_path":   img_path,
            "text_content": text_trunc,
        }

        # 3. Upsert: delete existing entry (no-op if absent), then add fresh
        try:
            collection.delete(ids=[chroma_id])
        except Exception:
            pass

        collection.add(
            ids=[chroma_id],
            embeddings=[embedding.tolist()],   # ChromaDB requires a Python list
            metadatas=[page_metadata],
            documents=[text_trunc],
        )

        embedded += 1
        torch.cuda.empty_cache()   # release VRAM after every forward pass

    logger.info(
        f"✅ ChromaDB: {embedded}/{len(pages)} pages stored for '{display_name}' "
        f"| Account {workspace.account_id} index total: {collection.count()} pages"
    )

    doc_meta = {
        "id":         doc_id,
        "name":       display_name,
        "type":       ext,
        "size":       file_path.stat().st_size,
        # A content hash, so the same file uploaded twice is recognised rather
        # than re-rendered, re-embedded and stored again.  See document_storage.
        "digest":     _file_digest(file_path),
        "pages":      pages,
        "page_count": len(pages),
        "status":     "indexed",
        "created_at": datetime.now().isoformat(),
    }
    meta.setdefault("documents", {})[doc_id] = doc_meta
    workspace.save_metadata(meta)
    return doc_meta


print("✅ Ingestion pipeline ready — ColPali embeddings stored in ChromaDB")


# ======================================================================
# CELL 6
# ======================================================================

"""OpenAI-compatible CLIProxyAPI client used by BuildMarshalAI.

The adapter intentionally accepts the message format previously consumed by
Qwen2.5-VL (``{"type": "image", "image": "/path/page.png"}``) so chat,
document generation, Gmail drafting, and Calendar interpretation can keep
calling the existing ``vl_generate`` interface.
"""


import base64
import mimetypes
import os
from pathlib import Path
from typing import Any, Mapping, Sequence

import requests


class CLIProxyError(RuntimeError):
    """Raised when CLIProxyAPI is unavailable or returns an invalid response."""


def _as_bool(value: str | None, default: bool = True) -> bool:
    if value is None:
        return default
    return value.strip().lower() not in {"0", "false", "no", "off"}


class CLIProxyClient:
    """Small synchronous client for CLIProxyAPI's OpenAI-compatible surface."""

    def __init__(
        self,
        base_url: str | None = None,
        api_key: str | None = None,
        preferred_model: str | None = None,
        timeout_seconds: int | None = None,
        verify_ssl: bool | None = None,
        session: requests.Session | None = None,
    ) -> None:
        self.base_url = (base_url or os.getenv("CLIPROXY_BASE_URL") or "http://127.0.0.1:8317/v1").rstrip("/")
        self.api_key = api_key if api_key is not None else os.getenv("CLIPROXY_API_KEY", "")
        self.preferred_model = preferred_model or os.getenv("CLIPROXY_MODEL", "gemini-3.7-flash")
        self.timeout_seconds = timeout_seconds or int(os.getenv("CLIPROXY_TIMEOUT_SECONDS", "240"))
        self.verify_ssl = _as_bool(os.getenv("CLIPROXY_VERIFY_SSL"), True) if verify_ssl is None else verify_ssl
        self.session = session or requests.Session()
        self._selected_model: str | None = None

    @property
    def headers(self) -> dict[str, str]:
        headers = {"Accept": "application/json", "Content-Type": "application/json"}
        if self.api_key:
            headers["Authorization"] = f"Bearer {self.api_key}"
        return headers

    def _request(
        self,
        method: str,
        path: str,
        request_timeout: int | tuple[int, int] | None = None,
        **kwargs: Any,
    ) -> requests.Response:
        url = f"{self.base_url}/{path.lstrip('/')}"
        try:
            response = self.session.request(
                method,
                url,
                headers=self.headers,
                timeout=request_timeout or self.timeout_seconds,
                verify=self.verify_ssl,
                **kwargs,
            )
        except requests.RequestException as exc:
            raise CLIProxyError(
                f"Cannot reach CLIProxyAPI at {self.base_url}. Start the proxy and check "
                "CLIPROXY_BASE_URL. A Kaggle notebook cannot use your PC's 127.0.0.1."
            ) from exc
        if not response.ok:
            detail = response.text.strip()[:1000]
            raise CLIProxyError(f"CLIProxyAPI returned HTTP {response.status_code}: {detail}")
        return response

    def list_models(self, request_timeout: int | tuple[int, int] | None = None) -> list[str]:
        payload = self._request("GET", "models", request_timeout=request_timeout).json()
        models = payload.get("data", []) if isinstance(payload, Mapping) else []
        return [str(item["id"]) for item in models if isinstance(item, Mapping) and item.get("id")]

    def select_model(self, requested: str | None = None, refresh: bool = False) -> str:
        """Select an exposed Antigravity model without trusting a stale alias.

        Frontend values left over from the old Qwen setup are intentionally
        ignored. If the preferred identifier is absent, the newest exposed
        Gemini Flash identifier is selected from ``/v1/models``.
        """
        if self._selected_model and not refresh and not requested:
            return self._selected_model

        requested = (requested or "").strip()
        if "qwen" in requested.lower():
            requested = ""
        models = self.list_models()
        if not models:
            raise CLIProxyError("CLIProxyAPI returned no models. Complete Antigravity OAuth login first.")

        preferred = requested or self.preferred_model
        if preferred in models:
            self._selected_model = preferred
            return preferred

        flash_models = [model for model in models if "gemini" in model.lower() and "flash" in model.lower()]
        if flash_models:
            self._selected_model = sorted(flash_models, reverse=True)[0]
            return self._selected_model

        raise CLIProxyError(
            f"Requested model '{preferred}' is unavailable and no Gemini Flash model was exposed. "
            f"Available models: {', '.join(models[:30])}"
        )

    @staticmethod
    def image_data_url(value: str | os.PathLike[str]) -> str:
        raw_value = str(value)
        if raw_value.startswith(("data:", "http://", "https://")):
            return raw_value
        path = Path(raw_value)
        if not path.is_file():
            raise CLIProxyError(f"Image does not exist: {path}")
        mime = mimetypes.guess_type(path.name)[0] or "image/png"
        encoded = base64.b64encode(path.read_bytes()).decode("ascii")
        return f"data:{mime};base64,{encoded}"

    def convert_messages(self, messages: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
        converted: list[dict[str, Any]] = []
        for message in messages:
            role = str(message.get("role", "user"))
            content = message.get("content", "")
            if not isinstance(content, list):
                converted.append({"role": role, "content": str(content)})
                continue

            parts: list[dict[str, Any]] = []
            for part in content:
                if not isinstance(part, Mapping):
                    parts.append({"type": "text", "text": str(part)})
                    continue
                part_type = str(part.get("type", "text"))
                if part_type in {"image", "image_url"}:
                    image_value: Any = part.get("image")
                    if image_value is None:
                        image_value = part.get("image_url")
                    if isinstance(image_value, Mapping):
                        image_value = image_value.get("url")
                    if image_value:
                        parts.append({
                            "type": "image_url",
                            "image_url": {"url": self.image_data_url(str(image_value))},
                        })
                else:
                    parts.append({"type": "text", "text": str(part.get("text", ""))})
            converted.append({"role": role, "content": parts})
        return converted

    @staticmethod
    def _extract_text(payload: Mapping[str, Any]) -> str:
        try:
            content = payload["choices"][0]["message"]["content"]
        except (KeyError, IndexError, TypeError) as exc:
            raise CLIProxyError(f"Unexpected CLIProxyAPI response: {str(payload)[:1000]}") from exc
        if isinstance(content, str):
            return content
        if isinstance(content, list):
            return "\n".join(
                str(item.get("text", ""))
                for item in content
                if isinstance(item, Mapping) and item.get("text")
            )
        return str(content)

    def chat(
        self,
        messages: Sequence[Mapping[str, Any]],
        max_tokens: int = 4096,
        model: str | None = None,
    ) -> tuple[str, str]:
        selected_model = self.select_model(model)
        response = self._request(
            "POST",
            "chat/completions",
            json={
                "model": selected_model,
                "messages": self.convert_messages(messages),
                "max_tokens": max_tokens,
                "temperature": 0,
                "stream": False,
            },
        )
        return self._extract_text(response.json()), selected_model

    def status(self) -> dict[str, Any]:
        try:
            models = self.list_models(request_timeout=(3, 10))
            selected = self.select_model()
            return {
                "connected": True,
                "base_url": self.base_url,
                "selected_model": selected,
                "model_count": len(models),
            }
        except Exception as exc:
            return {
                "connected": False,
                "base_url": self.base_url,
                "selected_model": None,
                "error": str(exc),
            }


CLIPROXY_CLIENT = CLIProxyClient()
CLIPROXY_LABEL = f"cliproxyapi:{CLIPROXY_CLIENT.preferred_model}"
_LAST_GENERATION_MODEL = None

def vl_generate(messages: list, max_new_tokens: int = 4096, model: str | None = None) -> str:
    """Compatibility entry point used by chat, document, Gmail, and Calendar flows."""
    global _LAST_GENERATION_MODEL, CLIPROXY_LABEL
    response, selected_model = CLIPROXY_CLIENT.chat(
        messages, max_tokens=max_new_tokens, model=model
    )
    _LAST_GENERATION_MODEL = selected_model
    CLIPROXY_LABEL = f"cliproxyapi:{selected_model}"
    return response

def image_to_base64(image_path: str) -> Optional[str]:
    try:
        with open(image_path, "rb") as handle:
            data = base64.b64encode(handle.read()).decode("utf-8")
        return f"data:image/png;base64,{data}"
    except Exception:
        return None

async def generate_response(
    query: str,
    context_pages: List[Dict],
    history: List[Dict],
    model: str | None = None,
) -> Dict:
    system_prompt = (
        "You are BuildMarshalAI, an intelligent construction document assistant.\n"
        "Answer only from the supplied project-document evidence. If the evidence does not "
        "contain the answer, say so. Cite document names and page numbers. Use clear markdown "
        "and describe relevant drawings or diagrams. Cite factual claims using the supplied "
        "numeric source labels, for example [1] or [2]. Never invent a citation number."
    )
    messages = [{"role": "system", "content": system_prompt}]
    for message in history[-6:]:
        messages.append({
            "role": message.get("role", "user"),
            "content": message.get("content", ""),
        })

    user_content = []
    images_added = 0
    for page in context_pages:
        image_path = page.get("image_path")
        if images_added >= 4:
            break
        if image_path and os.path.exists(image_path):
            user_content.append({"type": "image", "image": image_path})
            images_added += 1

    context_text = "Relevant document pages:\n"
    for index, page in enumerate(context_pages):
        context_text += (
            f"\n[Source {index + 1}: {page['doc_name']}, Page {page['page']}"
            f" | similarity={page.get('score', 0):.3f}]"
        )
        if page.get("text_content"):
            context_text += f"\n{page['text_content'][:800]}"
    user_content.append({"type": "text", "text": f"{context_text}\n\nUser question: {query}"})
    messages.append({"role": "user", "content": user_content})

    try:
        response = await asyncio.get_running_loop().run_in_executor(
            None, lambda: vl_generate(messages, max_new_tokens=4096, model=model)
        )
    except CLIProxyError as exc:
        logger.error("CLIProxyAPI inference error: %s", exc)
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    except Exception as exc:
        logger.exception("Generation failed")
        raise HTTPException(status_code=500, detail=f"Generation failed: {exc}") from exc

    sources = []
    for rank, page in enumerate(context_pages, 1):
        source = {
            "doc_name": page["doc_name"], "doc_id": page["doc_id"],
            "page": page["page"], "score": page.get("score", 0.0), "image_url": None,
            "rank": rank,
            "query": query,
            "retrieval_method": page.get("retrieval_method", "hybrid retrieval"),
            "score_breakdown": page.get("score_breakdown", {}),
            "matched_terms": page.get("matched_terms", []),
            "source_type": page.get("source_type", "project_document"),
            "evidence_text": page.get("text_content", "")[:1600],
        }
        if page.get("image_path"):
            source["image_url"] = image_to_base64(page["image_path"])
        sources.append(source)
    return {"response": response, "sources": sources, "model_used": _LAST_GENERATION_MODEL}

print("Query pipeline ready - ColPali retrieval + CLIProxyAPI Antigravity generation")


# ======================================================================
# CELL 7
# ======================================================================

# =======================================================================
# CELL 8 — FastAPI application
# =======================================================================

app = FastAPI(
    title="BuildMarshalAI Backend",
    description="ColPali (per-account ChromaDB vector stores) + CLIProxyAPI Antigravity generation",
    version="3.0.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# =======================================================================
# Account resolution
# =======================================================================
#
# The account system is registered with the other integration modules, after
# this cell has created ``app``.  Routes therefore depend on these two thin
# wrappers, which delegate to the resolver published at registration time.
# Until that happens every account-scoped route answers 503 rather than
# silently serving data with no tenant attached.
# =======================================================================

ACCOUNT_RESOLVER = None


async def require_account(request: Request):
    """The authenticated caller's AccountContext, or 401/503."""
    if ACCOUNT_RESOLVER is None:
        raise HTTPException(503, "The account system is still starting up")
    context = ACCOUNT_RESOLVER(request.headers.get("Authorization", "").removeprefix("Bearer ").strip())
    if context is None:
        raise HTTPException(401, "Sign in to continue")
    return context


async def optional_account(request: Request):
    """The caller's AccountContext when signed in, otherwise None."""
    if ACCOUNT_RESOLVER is None:
        return None
    return ACCOUNT_RESOLVER(request.headers.get("Authorization", "").removeprefix("Bearer ").strip())


def sanitize_doc_id(value: str | None) -> str:
    """Client-supplied document ids become file names, so constrain them."""
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "", str(value or "")).strip("._-")
    if not cleaned or cleaned in {".", ".."}:
        return uuid.uuid4().hex[:12]
    return cleaned[:64]


# ── Public routes ─────────────────────────────────────────────────────────────

@app.get("/api/health")
async def health_check(context=Depends(optional_account)):
    gpu_info = []
    if torch.cuda.is_available():
        for i in range(NUM_GPUS):
            free, total = torch.cuda.mem_get_info(i)
            hosts = []
            if i == COLPALI_DEVICE_IDX:
                hosts.append("colpali")
            gpu_info.append({
                "index":    i,
                "name":     torch.cuda.get_device_name(i),
                "free_gb":  round(free / 1e9, 2),
                "total_gb": round(total / 1e9, 2),
                "hosts":    hosts,
            })
    payload = {
        "status":          "healthy",
        "gpu_available":   torch.cuda.is_available(),
        "gpu_count":       NUM_GPUS,
        "gpus":            gpu_info,
        "authenticated":   context is not None,
        "generator": CLIPROXY_LABEL,
        "cliproxy_base_url": CLIPROXY_CLIENT.base_url,
    }
    # Corpus sizes describe one account, so they are reported only to a caller
    # that has one.
    if context is not None:
        payload["chroma_pages"] = context.workspace.collection.count()
        payload["documents_count"] = len(context.workspace.load_metadata().get("documents", {}))
        payload["account_id"] = context.account_id
    else:
        payload["chroma_pages"] = 0
        payload["documents_count"] = 0
    return payload


@app.get("/api/cliproxy/status")
async def cliproxy_status():
    return await asyncio.get_running_loop().run_in_executor(None, CLIPROXY_CLIENT.status)


# ── Account-scoped routes ─────────────────────────────────────────────────────

@app.get("/api/status")
async def get_status(context=Depends(require_account)):
    workspace = context.workspace
    docs  = workspace.load_metadata().get("documents", {})
    total = sum(d.get("page_count", 0) for d in docs.values())
    return {
        "account_id":      context.account_id,
        "documents_count": len(docs),
        "total_pages":     total,
        "chroma_pages":    workspace.collection.count(),
        "retriever":       f"colpali-v1.2 + chromadb ({COLPALI_DEVICE})",
        "generator":       CLIPROXY_LABEL,
        "cliproxy_base_url": CLIPROXY_CLIENT.base_url,
        "gpu_count":       NUM_GPUS,
    }


@app.post("/api/upload")
async def upload_document(
    file: UploadFile = File(...),
    doc_id: str = Form(None),
    context=Depends(require_account),
):
    workspace = context.workspace
    doc_id = sanitize_doc_id(doc_id)
    ext    = Path(file.filename or "").suffix.lower()

    allowed = {
    # Documents
    '.pdf', '.xlsx', '.xls', '.csv', '.doc', '.docx', '.pptx', '.ppt',
    '.txt', '.rtf', '.odt', '.ods',
    # Images
    '.png', '.jpg', '.jpeg', '.gif', '.webp', '.bmp', '.tiff',
    # Audio
    '.mp3', '.wav', '.ogg', '.m4a', '.flac', '.aac', '.wma', '.opus',
    # Video
    '.mp4', '.webm', '.mov', '.avi', '.mkv',
    # Archives & Code
    '.zip', '.rar', '.7z', '.tar', '.gz',
    '.json', '.xml', '.html', '.css', '.js', '.py', '.java', '.c', '.cpp', '.md',
}

    if ext not in allowed:
        raise HTTPException(400, detail=f"Unsupported file type: {ext}")
    save_path = workspace.docs_dir / f"{doc_id}{ext}"
    # Stream to disk so a large upload cannot exhaust memory, and stop as soon
    # as the ceiling is crossed rather than after the whole body has arrived.
    written = 0
    try:
        with save_path.open("wb") as destination:
            while chunk := await file.read(1024 * 1024):
                written += len(chunk)
                if written > MAX_UPLOAD_BYTES:
                    destination.close()
                    save_path.unlink(missing_ok=True)
                    raise HTTPException(
                        413, detail=f"File exceeds the {MAX_UPLOAD_BYTES // (1024 * 1024)} MB upload limit"
                    )
                destination.write(chunk)
    except HTTPException:
        raise
    except Exception as e:
        save_path.unlink(missing_ok=True)
        raise HTTPException(500, detail=f"Could not save file: {e}")
    # Indexing is the expensive part -- rendering every page, embedding each one,
    # and storing the result.  The same bytes already indexed cost all of that
    # again for nothing, so look before doing the work.
    try:
        digest = _file_digest(save_path)
        existing = next(
            (dict(stored, id=stored_id)
             for stored_id, stored in workspace.load_metadata().get("documents", {}).items()
             if stored_id != doc_id and stored.get("digest") == digest),
            None,
        )
    except Exception:
        existing = None
    if existing:
        save_path.unlink(missing_ok=True)
        return {
            "id":      existing["id"],
            "name":    existing.get("name", file.filename),
            "pages":   existing.get("page_count", 0),
            "status":  "duplicate",
            "message": (
                f"Already indexed as \"{existing.get('name', existing['id'])}\"; "
                "the existing copy was reused."
            ),
        }

    try:
        doc_meta = ingest_document(save_path, doc_id, workspace, display_name=file.filename)
        return {
            "id":      doc_id,
            "name":    file.filename,
            "pages":   doc_meta["page_count"],
            "status":  "indexed",
            "message": (
                f"Indexed {doc_meta['page_count']} pages "
                f"(account index total: {workspace.collection.count()})"
            ),
        }
    except Exception as e:
        logger.error(f"Ingestion error: {e}")
        return {"id": doc_id, "name": file.filename, "pages": 0,
                "status": "error", "message": str(e)}


@app.get("/api/documents")
async def list_documents(context=Depends(require_account)):
    meta = context.workspace.load_metadata()
    return {"documents": [
        {
            "id":         did,
            "name":       d["name"],
            "type":       d["type"],
            "size":       d.get("size", 0),
            "pages":      d.get("page_count", 0),
            "status":     d.get("status", "indexed"),
            "created_at": d.get("created_at"),
        }
        for did, d in meta.get("documents", {}).items()
    ]}


@app.get("/api/documents/{doc_id}/status")
async def document_status(doc_id: str, context=Depends(require_account)):
    doc = context.workspace.load_metadata().get("documents", {}).get(doc_id)
    if not doc:
        raise HTTPException(404, detail="Document not found")
    return {
        "id":     doc_id,
        "status": doc.get("status", "unknown"),
        "pages":  doc.get("page_count", 0),
    }


def delete_document_files(workspace, doc_id: str) -> None:
    """Remove every trace of one document.

    Named and reachable so the storage sweep can collapse a duplicate exactly
    the way a hand-deletion does, rather than growing a second, divergent
    cleanup path.
    """
    meta = workspace.load_metadata()
    doc = meta.get("documents", {}).get(doc_id)
    if not doc:
        raise HTTPException(404, detail="Document not found")

    # 1. Remove page screenshots from this account's workspace.  Paths are
    #    resolved through the workspace so a metadata entry can never name a
    #    file outside it.
    for page in doc.get("pages", []):
        resolved = workspace.resolve_page_path(page.get("image_path"))
        if resolved:
            Path(resolved).unlink(missing_ok=True)
    #    A prefix glob would also match a longer id that starts with this one,
    #    so match the stem exactly and let the suffix vary.
    for stored in workspace.docs_dir.iterdir():
        if stored.is_file() and stored.stem == doc_id:
            stored.unlink(missing_ok=True)

    # 2. Remove this document's vectors from the account's collection.
    try:
        collection = workspace.collection
        collection.delete(where={"doc_id": doc_id})
        logger.info(
            f"Vector index: removed entries for doc_id={doc_id} "
            f"| account {workspace.account_id} now has {collection.count()} pages"
        )
    except Exception as e:
        logger.warning(f"Vector index delete failed for {doc_id}: {e}")

    # 3. The ColPali multi-vector cache and the vision tiles are keyed by page
    #    filename, not by doc_id, so they outlive the document unless they are
    #    removed here.  Left behind they are unreachable and permanent.
    import re as _re

    for page in doc.get("pages", []):
        stored = page.get("image_path") or ""
        if not stored:
            continue
        stem = _re.sub(r"[^A-Za-z0-9._-]+", "_", Path(stored).stem)
        (Path(workspace.multivector_dir) / f"{stem}.npy").unlink(missing_ok=True)
        (Path(workspace.vision_cache_dir) / f"{stem}_512.jpg").unlink(missing_ok=True)

    del meta["documents"][doc_id]
    workspace.save_metadata(meta)


@app.delete("/api/documents/{doc_id}")
async def delete_document(doc_id: str, context=Depends(require_account)):
    delete_document_files(context.workspace, doc_id)
    return {"message": "Document deleted", "id": doc_id}


@app.post("/api/chat")
async def chat(request: Request, context=Depends(require_account)):
    body  = await request.json()
    query = body.get("query", "").strip()
    if not query:
        raise HTTPException(400, detail="Query cannot be empty")

    workspace = context.workspace
    settings = workspace.load_settings()
    history = body.get("history", [])
    model   = body.get("model") or settings.get("model")
    top_k   = min(int(body.get("top_k") or settings.get("top_k", 5)), 20)
    project_id = body.get("project_id")
    if project_id and project_id not in workspace.load_projects():
        raise HTTPException(404, detail="Project not found")

    logger.info(f"Query [{context.account_id}]: {query[:80]}...")

    # Step 1: Retrieve top-k pages from this account's index only
    context_pages = retrieve_context(query, top_k=top_k, project_id=project_id, workspace=workspace)
    logger.info(f"Retrieved {len(context_pages)} pages for account {context.account_id}")

    # Step 2: Generate through CLIProxyAPI using text plus retrieved page images
    return await generate_response(query, context_pages, history, model=model)


@app.get("/api/pages/{doc_id}/{page_num}")
async def get_page_image(doc_id: str, page_num: int, context=Depends(require_account)):
    workspace = context.workspace
    doc = workspace.load_metadata().get("documents", {}).get(doc_id)
    if not doc:
        raise HTTPException(404, detail="Document not found")
    for page in doc.get("pages", []):
        if page.get("page_num") == page_num:
            img_path = workspace.resolve_page_path(page.get("image_path"))
            if img_path:
                # Detect content type from extension
                ext = Path(img_path).suffix.lower()
                content_types = {
                    '.png': 'image/png', '.jpg': 'image/jpeg', '.jpeg': 'image/jpeg',
                    '.gif': 'image/gif', '.webp': 'image/webp',
                    '.mp3': 'audio/mpeg', '.wav': 'audio/wav', '.ogg': 'audio/ogg',
                    '.m4a': 'audio/mp4', '.flac': 'audio/flac', '.aac': 'audio/aac',
                    '.mp4': 'video/mp4', '.webm': 'video/webm', '.mov': 'video/quicktime',
                }
                ct = content_types.get(ext, 'application/octet-stream')
                # Somebody looked at this document: that is what separates a
                # document worth keeping cached from one indexed and forgotten.
                try:
                    meta = workspace.load_metadata()
                    if doc_id in meta.get("documents", {}):
                        meta["documents"][doc_id]["accessed_at"] = datetime.now(timezone.utc).isoformat()
                        workspace.save_metadata(meta)
                except Exception:
                    pass
                with open(img_path, "rb") as f:
                    return StreamingResponse(BytesIO(f.read()), media_type=ct)
    raise HTTPException(404, detail="Page not found")


@app.get("/api/chroma/stats")
async def chroma_stats(context=Depends(require_account)):
    """Inspect this account's vector collection: total count + sample entries."""
    collection = context.workspace.collection
    count  = collection.count()
    sample = []
    if count > 0:
        peek = collection.peek(limit=5)
        for i, meta in enumerate(peek.get("metadatas", [])):
            sample.append({
                "id":       peek["ids"][i],
                "doc_name": meta.get("doc_name"),
                "page_num": meta.get("page_num"),
                "doc_id":   meta.get("doc_id"),
            })
    return {"total_pages": count, "sample": sample}


print("FastAPI app defined — account-scoped endpoints ready")


# ======================================================================
# CELL 8
# ======================================================================

from fpdf import FPDF
from io import BytesIO
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import List, Optional

# Models for the API payload
class ChatMessage(BaseModel):
    role: str
    content: str
    timestamp: Optional[str] = None

class ChatExportRequest(BaseModel):
    title: str
    messages: List[ChatMessage]


def sanitize_for_pdf(text: str) -> str:
    if not text: return ""
    return str(text).encode("latin-1", "replace").decode("latin-1")


# Endpoint 1: Project overview PDF, built from the caller's own project record
@app.post("/api/projects/{project_id}/generate-document")
async def generate_project_overview(project_id: str, context=Depends(require_account)):
    project = context.workspace.load_projects().get(project_id)
    if not project:
        raise HTTPException(404, detail="Project not found")

    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    # Title
    pdf.set_font("helvetica", "B", 24)
    pdf.cell(0, 10, "Project Overview Report", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(10)

    # Project Info
    pdf.set_font("helvetica", "B", 14)
    pdf.cell(0, 10, sanitize_for_pdf(f"Project Name: {project.get('name', 'N/A')}"), new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("helvetica", "", 12)
    for label, key in (
        ("Project Code", "project_code"), ("Manager", "manager"), ("Status", "status"),
        ("Start Date", "start_date"), ("End Date", "end_date"),
    ):
        pdf.cell(0, 8, sanitize_for_pdf(f"{label}: {project.get(key) or 'N/A'}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(10)

    pdf.set_font("helvetica", "I", 10)
    pdf.cell(0, 10, f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M')}", new_x="LMARGIN", new_y="NEXT")

    pdf_buffer = BytesIO()
    pdf.output(pdf_buffer)
    pdf_buffer.seek(0)

    filename = f"Project_{project_id}_Report.pdf"
    return StreamingResponse(
        pdf_buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# Endpoint 2: Export Chat as PDF
@app.post("/api/chat/export-pdf")
def export_chat_pdf(req: ChatExportRequest, context=Depends(require_account)):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    # Title
    pdf.set_font("helvetica", "B", 18)
    pdf.cell(0, 10, sanitize_for_pdf(f"Chat Session: {req.title}"), new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(5)

    for msg in req.messages:
        role = "Marshal AI" if msg.role == "bot" else "User"
        pdf.set_font("helvetica", "B", 12)

        time_str = f" [{msg.timestamp}]" if msg.timestamp else ""
        pdf.cell(0, 8, sanitize_for_pdf(f"{role}{time_str}"), new_x="LMARGIN", new_y="NEXT")

        pdf.set_font("helvetica", "", 11)
        # Handle multi-line content
        pdf.multi_cell(0, 6, sanitize_for_pdf(msg.content))
        pdf.ln(4)

    pdf_buffer = BytesIO()
    pdf.output(pdf_buffer)
    pdf_buffer.seek(0)

    filename = "Marshal_Chat_Export.pdf"
    return StreamingResponse(
        pdf_buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# ======================================================================
# CELL 9
# ======================================================================

# =============================================================================
# User management and authentication
#
# These routes moved into ``backend/accounts.py`` and are registered with the
# other integration modules further down.  The module owns:
#
#   POST   /api/auth/register     create an account plus its owner login
#   POST   /api/auth/login        exchange credentials for a session token
#   POST   /api/auth/logout       revoke the current session
#   GET    /api/auth/me           the signed-in user, account, and settings
#   PUT    /api/auth/profile      update the signed-in user's own profile
#   PUT    /api/auth/password     rotate a password and re-issue the session
#   GET    /api/account           account summary and per-account usage
#   PUT    /api/account           rename the account (administrators)
#   DELETE /api/account           delete the account and all its data (owner)
#   GET    /PUT /api/account/settings
#   GET    /PUT /api/conversations, DELETE /api/conversations/{id}
#   GET/POST/PUT/DELETE /api/users[/{user_id}]   members of the caller's account
#
# Passwords are PBKDF2-HMAC-SHA256 with a per-user salt.  Sessions are random
# bearer tokens stored as digests with an absolute and an idle expiry, so a
# leaked store cannot be replayed and a user id is no longer a credential.
# =============================================================================


# ======================================================================
# CELL 10
# ======================================================================

# =============================================================================
# Projects + Tasks
#
# Projects and tasks are stored in the caller's workspace (projects.json and
# tasks.json) rather than in process memory, so they survive a restart and stay
# consistent with the documents that reference them by project_id.
# =============================================================================

def _make_project(data: dict) -> dict:
    return {
        "id":           str(uuid.uuid4()),
        "name":         data.get("name", "").strip(),
        "project_code": data.get("project_code", "").strip(),
        "manager":      data.get("manager", "").strip(),
        "type":         data.get("type", "").strip(),
        "status":       data.get("status", "Active").strip(),
        "start_date":   data.get("start_date", ""),
        "end_date":     data.get("end_date", ""),
        "description":  data.get("description", "").strip(),
        "address_line1":data.get("address_line1", "").strip(),
        "address_line2":data.get("address_line2", "").strip(),
        "city":         data.get("city", "").strip(),
        "state":        data.get("state", "").strip(),
        "postal_code":  data.get("postal_code", "").strip(),
        "country":      data.get("country", "").strip(),
        "archived":     False,
        "created_at":   datetime.now(timezone.utc).isoformat(),
    }

def _project_or_404(workspace, project_id: str) -> dict:
    project = workspace.load_projects().get(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    return project


# ── Project Routes ─────────────────────────────────────────────────────────────

@app.get("/api/projects")
async def list_projects(
    name:          str = "",
    manager:       str = "",
    type:          str = "",        # comma-separated values
    status:        str = "",        # comma-separated values
    show_archived: bool = False,
    start_after:   str = "",
    start_before:  str = "",
    end_after:     str = "",
    end_before:    str = "",
    page:          int = 1,
    per_page:      int = 10,
    context=Depends(require_account),
):
    projects = list(context.workspace.load_projects().values())

    # Archived filter
    if not show_archived:
        projects = [p for p in projects if not p.get("archived")]

    # Text filters
    if name and len(name) >= 2:
        projects = [p for p in projects if name.lower() in p["name"].lower()]
    if manager and len(manager) >= 2:
        projects = [p for p in projects if manager.lower() in p["manager"].lower()]

    # Multi-value filters (comma-separated)
    if type:
        types = [t.strip() for t in type.split(",") if t.strip()]
        if types:
            projects = [p for p in projects if p["type"] in types]
    if status:
        statuses = [s.strip() for s in status.split(",") if s.strip()]
        if statuses:
            projects = [p for p in projects if p["status"] in statuses]

    # Date range filters — skip projects that have no date set
    if start_after:
        projects = [p for p in projects if p.get("start_date") and p["start_date"] >= start_after]
    if start_before:
        projects = [p for p in projects if p.get("start_date") and p["start_date"] <= start_before]
    if end_after:
        projects = [p for p in projects if p.get("end_date") and p["end_date"] >= end_after]
    if end_before:
        projects = [p for p in projects if p.get("end_date") and p["end_date"] <= end_before]

    projects.sort(key=lambda p: str(p.get("created_at", "")), reverse=True)
    per_page = max(1, min(per_page, 100))
    total = len(projects)
    pages = max(1, (total + per_page - 1) // per_page)
    page  = max(1, min(page, pages))
    start = (page - 1) * per_page
    return {
        "projects": projects[start: start + per_page],
        "total": total, "page": page, "per_page": per_page, "pages": pages,
    }


@app.get("/api/projects/{project_id}")
async def get_project(project_id: str, context=Depends(require_account)):
    return _project_or_404(context.workspace, project_id)


@app.post("/api/projects")
async def create_project(req: Request, context=Depends(require_account)):
    context.require("project.create", "creating projects")
    data = await req.json()
    if not data.get("name", "").strip():
        raise HTTPException(status_code=422, detail="Name is required")
    if not data.get("project_code", "").strip():
        raise HTTPException(status_code=422, detail="Project code is required")
    workspace = context.workspace
    projects = workspace.load_projects()
    # Project codes are unique within an account, not across the install.
    for p in projects.values():
        if p["project_code"].upper() == data["project_code"].strip().upper():
            raise HTTPException(status_code=409, detail="Project code already in use")
    project = _make_project(data)
    projects[project["id"]] = project
    workspace.save_projects(projects)
    tasks = workspace.load_tasks()
    tasks[project["id"]] = []
    workspace.save_tasks(tasks)
    return project


@app.put("/api/projects/{project_id}")
async def update_project(project_id: str, req: Request, context=Depends(require_account)):
    context.require("project.update", "editing projects")
    workspace = context.workspace
    projects = workspace.load_projects()
    p = projects.get(project_id)
    if not p:
        raise HTTPException(status_code=404, detail="Project not found")
    data = await req.json()
    for field in ["name", "project_code", "manager", "type", "status",
                  "start_date", "end_date", "description",
                  "address_line1", "address_line2", "city", "state",
                  "postal_code", "country", "archived"]:
        if field in data:
            p[field] = data[field]
    projects[project_id] = p
    workspace.save_projects(projects)
    return p


@app.delete("/api/projects/{project_id}")
async def delete_project(project_id: str, context=Depends(require_account)):
    workspace = context.workspace
    projects = workspace.load_projects()
    if project_id not in projects:
        raise HTTPException(status_code=404, detail="Project not found")
    del projects[project_id]
    workspace.save_projects(projects)
    tasks = workspace.load_tasks()
    tasks.pop(project_id, None)
    workspace.save_tasks(tasks)
    # Procurement lines belong to the project too; leaving them would orphan
    # rows that nothing can reach or clean up.
    procurement = workspace.load_procurement()
    if procurement.pop(project_id, None) is not None:
        workspace.save_procurement(procurement)
    # Documents keep their content but lose the link to a project that is gone,
    # so /api/documents and retrieval stay consistent.
    metadata = workspace.load_metadata()
    detached = False
    for document in metadata.get("documents", {}).values():
        if document.get("project_id") == project_id:
            document["project_id"] = None
            detached = True
    if detached:
        workspace.save_metadata(metadata)
    return {"message": "Project deleted", "id": project_id}


# Task routes are registered from backend/tasks.py with the other
# integration modules further down.


# ======================================================================
# CELL 11
# ======================================================================


# =======================================================================
# Management APIs (Trades, Vendors, Team Members)
#
# Each account gets its own management.json, seeded with the standard trade
# list the first time it is read.  A new workspace therefore starts with the
# same defaults the product has always shipped, but nothing is shared.
# =======================================================================

# ── Trades CRUD ──
@app.get("/api/trades")
async def list_trades(context=Depends(require_account)):
    return {"trades": context.workspace.load_mgmt().get("trades", [])}

@app.post("/api/trades")
async def create_trade(request: Request, context=Depends(require_account)):
    body = await request.json()
    mgmt = context.workspace.load_mgmt()
    trade = {
        "id": f"tr-{uuid.uuid4().hex[:8]}",
        "name": body.get("name", "").strip(),
        "description": body.get("description", "-"),
        "status": body.get("status", "Active")
    }
    if not trade["name"]:
        raise HTTPException(400, detail="Name is required")
    mgmt["trades"].append(trade)
    context.workspace.save_mgmt(mgmt)
    return trade

@app.put("/api/trades/{trade_id}")
async def update_trade(trade_id: str, request: Request, context=Depends(require_account)):
    body = await request.json()
    mgmt = context.workspace.load_mgmt()
    for t in mgmt["trades"]:
        if t["id"] == trade_id:
            t["name"] = body.get("name", t["name"])
            t["description"] = body.get("description", t["description"])
            t["status"] = body.get("status", t["status"])
            context.workspace.save_mgmt(mgmt)
            return t
    raise HTTPException(404, detail="Trade not found")

@app.delete("/api/trades/{trade_id}")
async def delete_trade(trade_id: str, context=Depends(require_account)):
    mgmt = context.workspace.load_mgmt()
    remaining = [t for t in mgmt["trades"] if t["id"] != trade_id]
    if len(remaining) == len(mgmt["trades"]):
        raise HTTPException(404, detail="Trade not found")
    mgmt["trades"] = remaining
    context.workspace.save_mgmt(mgmt)
    return {"message": "Trade deleted", "id": trade_id}



# ── Vendors CRUD ──
@app.get("/api/vendors")
async def list_vendors(context=Depends(require_account)):
    return {"vendors": context.workspace.load_mgmt().get("vendors", [])}

@app.post("/api/vendors")
async def create_vendor(request: Request, context=Depends(require_account)):
    body = await request.json()
    mgmt = context.workspace.load_mgmt()
    vendor = {
        "id": f"v-{uuid.uuid4().hex[:8]}",
        "name": body.get("name", "").strip(),
        "vendorType": body.get("vendorType", "Material Supplier"),
        "trade": body.get("trade", ""),
        "activeProjects": body.get("activeProjects", 0),
        "status": body.get("status", "Active")
    }
    if not vendor["name"]:
        raise HTTPException(400, detail="Name is required")
    mgmt["vendors"].append(vendor)
    context.workspace.save_mgmt(mgmt)
    return vendor

@app.put("/api/vendors/{vendor_id}")
async def update_vendor(vendor_id: str, request: Request, context=Depends(require_account)):
    body = await request.json()
    mgmt = context.workspace.load_mgmt()
    for v in mgmt["vendors"]:
        if v["id"] == vendor_id:
            v["name"] = body.get("name", v["name"])
            v["vendorType"] = body.get("vendorType", v["vendorType"])
            v["trade"] = body.get("trade", v["trade"])
            v["status"] = body.get("status", v["status"])
            context.workspace.save_mgmt(mgmt)
            return v
    raise HTTPException(404, detail="Vendor not found")

@app.delete("/api/vendors/{vendor_id}")
async def delete_vendor(vendor_id: str, context=Depends(require_account)):
    mgmt = context.workspace.load_mgmt()
    remaining = [v for v in mgmt["vendors"] if v["id"] != vendor_id]
    if len(remaining) == len(mgmt["vendors"]):
        raise HTTPException(404, detail="Vendor not found")
    mgmt["vendors"] = remaining
    context.workspace.save_mgmt(mgmt)
    return {"message": "Vendor deleted", "id": vendor_id}


# ── Team Members CRUD ──
@app.get("/api/team-members")
async def list_team_members(context=Depends(require_account)):
    return {"team_members": context.workspace.load_mgmt().get("team_members", [])}

@app.post("/api/team-members")
async def create_team_member(request: Request, context=Depends(require_account)):
    body = await request.json()
    mgmt = context.workspace.load_mgmt()
    member = {
        "id": f"tm-{uuid.uuid4().hex[:8]}",
        "name": body.get("name", "").strip(),
        "email": body.get("email", ""),
        "department": body.get("department", "—"),
        "category": body.get("category", "internal"),
        "company": body.get("company", ""),
        "contactName": body.get("contactName", ""),
    }
    if not member["name"]:
        raise HTTPException(400, detail="Name is required")
    mgmt["team_members"].append(member)
    context.workspace.save_mgmt(mgmt)
    return member

@app.put("/api/team-members/{member_id}")
async def update_team_member(member_id: str, request: Request, context=Depends(require_account)):
    body = await request.json()
    mgmt = context.workspace.load_mgmt()
    for m in mgmt["team_members"]:
        if m["id"] == member_id:
            for key in ["name", "email", "department", "category", "company", "contactName"]:
                if key in body:
                    m[key] = body[key]
            context.workspace.save_mgmt(mgmt)
            return m
    raise HTTPException(404, detail="Team member not found")

@app.delete("/api/team-members/{member_id}")
async def delete_team_member(member_id: str, context=Depends(require_account)):
    mgmt = context.workspace.load_mgmt()
    remaining = [m for m in mgmt["team_members"] if m["id"] != member_id]
    if len(remaining) == len(mgmt["team_members"]):
        raise HTTPException(404, detail="Team member not found")
    mgmt["team_members"] = remaining
    context.workspace.save_mgmt(mgmt)
    return {"message": "Team member deleted", "id": member_id}


# ======================================================================
# CELL 12
# ======================================================================

# CELL 11b - Document generation + Google Workspace integration
import importlib.util, os, subprocess, sys, urllib.request
from pathlib import Path

_integration_modules = ("reportlab", "pypdf", "googleapiclient", "google.auth", "cryptography")
if any(importlib.util.find_spec(name) is None for name in _integration_modules):
    subprocess.run([
        sys.executable, "-m", "pip", "install", "-q",
        "reportlab==4.4.3", "pypdf==5.7.0",
        "google-api-python-client==2.176.0", "google-auth==2.40.3",
        "cryptography==45.0.5",
    ], check=True)
else:
    print("Document-generation and Google integration packages already installed")

if IS_KAGGLE:
    try:
        from kaggle_secrets import UserSecretsClient
        _secrets = UserSecretsClient()
        for _name in (
            "GOOGLE_CLIENT_ID", "GOOGLE_CLIENT_SECRET",
            "GOOGLE_TOKEN_ENCRYPTION_KEY", "GOOGLE_ALLOWED_ORIGINS",
            "MICROSOFT_CLIENT_ID", "MICROSOFT_CLIENT_SECRET",
            "MICROSOFT_TENANT_ID", "MICROSOFT_ALLOWED_ORIGINS",
        ):
            try:
                _value = _secrets.get_secret(_name)
                if _value:
                    os.environ[_name] = _value
            except Exception:
                pass
    except Exception as exc:
        print(f"Kaggle Secrets unavailable: {exc}")

if IS_KAGGLE:
    integration_dir = RUNTIME_ROOT / "integration"
    integration_dir.mkdir(parents=True, exist_ok=True)
    (integration_dir / "hybrid_retrieval.py").write_text(
        "\"\"\"Hybrid text + ColPali late-interaction retrieval for BuildMarshalAI.\"\"\"\n\nfrom __future__ import annotations\n\nimport logging\nimport os\nimport re\nfrom collections import Counter\nfrom pathlib import Path\nfrom typing import Any, Mapping\n\nimport numpy as np\nimport torch\nfrom PIL import Image\n\n\nLOGGER = logging.getLogger(\"BuildMarshalAI.retrieval\")\n\nSTOPWORDS = {\n    \"a\", \"an\", \"and\", \"are\", \"about\", \"can\", \"could\", \"details\", \"do\", \"for\",\n    \"from\", \"give\", \"i\", \"in\", \"is\", \"me\", \"of\", \"on\", \"please\", \"show\", \"tell\",\n    \"the\", \"this\", \"to\", \"very\", \"what\", \"with\", \"you\",\n}\n\n\ndef search_tokens(value: str) -> list[str]:\n    return [\n        token\n        for token in re.findall(r\"[a-z0-9]+\", str(value).casefold())\n        if token not in STOPWORDS\n    ]\n\n\ndef lexical_score(query: str, page: Mapping[str, Any]) -> float:\n    query_tokens = search_tokens(query)\n    if not query_tokens:\n        return 0.0\n    haystack = f\"{page.get('doc_name', '')} {page.get('text_content', '')}\".casefold()\n    haystack_tokens = search_tokens(haystack)\n    if not haystack_tokens:\n        return 0.0\n    counts = Counter(haystack_tokens)\n    unique_query = list(dict.fromkeys(query_tokens))\n    matched = sum(1 for token in unique_query if counts.get(token, 0))\n    coverage = matched / max(len(unique_query), 1)\n    frequency = sum(min(counts.get(token, 0), 4) for token in unique_query) / max(len(unique_query), 1)\n    phrase = \" \".join(unique_query)\n    phrase_bonus = 4.0 if len(unique_query) >= 2 and phrase in \" \".join(haystack_tokens) else 0.0\n    detail_intent = bool(re.search(\n        r\"\\b(detail|details|describe|description|information|spec|specification|specifications)\\b\",\n        query.casefold(),\n    ))\n    index_page = any(marker in haystack for marker in (\n        \"item index\", \"drawing index\", \"table of contents\", \"sheet index\",\n    )) or \"index\" in search_tokens(str(page.get(\"doc_name\", \"\")))\n    # An index often contains the exact requested label but only points to the\n    # real answer. For detail/specification questions, prefer the richer page.\n    richness_bonus = (\n        min(2.0, len(haystack_tokens) / 400.0)\n        if detail_intent and phrase_bonus\n        else 0.0\n    )\n    index_penalty = 3.0 if detail_intent and index_page else 0.0\n    return max(\n        0.0,\n        phrase_bonus + 3.0 * coverage + 0.35 * frequency\n        + richness_bonus - index_penalty,\n    )\n\n\ndef normalize_scores(values: Mapping[str, float]) -> dict[str, float]:\n    if not values:\n        return {}\n    low, high = min(values.values()), max(values.values())\n    if high - low < 1e-9:\n        return {key: (1.0 if high > 0 else 0.0) for key in values}\n    return {key: (value - low) / (high - low) for key, value in values.items()}\n\n\nclass HybridColPaliRetriever:\n    \"\"\"Retrieve candidates lexically and visually, then apply ColPali MaxSim.\"\"\"\n\n    def __init__(self, namespace: Mapping[str, Any]):\n        # The ColPali model and processor are stateless and shared.  Everything\n        # that holds data -- the vector collection, the page metadata, and the\n        # multi-vector cache -- comes from the workspace passed per call, so one\n        # retriever instance can never mix two accounts' pages.\n        self.model = namespace[\"COLPALI_MODEL\"]\n        self.processor = namespace[\"COLPALI_PROCESSOR\"]\n        self.device = namespace.get(\"COLPALI_DEVICE\", \"cuda:0\")\n\n    def _cache_path(self, image_path: str, workspace: Any) -> Path:\n        safe_stem = re.sub(r\"[^A-Za-z0-9._-]+\", \"_\", Path(image_path).stem)\n        cache_dir = Path(workspace.multivector_dir)\n        cache_dir.mkdir(parents=True, exist_ok=True)\n        return cache_dir / f\"{safe_stem}.npy\"\n\n    @torch.no_grad()\n    def _encode_image(self, image_path: str) -> torch.Tensor | None:\n        try:\n            with Image.open(image_path) as source:\n                image = source.convert(\"RGB\")\n            batch = self.processor.process_images([image]).to(self.model.device)\n            return self.model(**batch).squeeze(0).to(torch.float16).cpu()\n        except Exception as exc:\n            LOGGER.error(\"ColPali page encoding failed for %s: %s\", image_path, exc)\n            return None\n\n    def _page_multivector(self, image_path: str, workspace: Any) -> torch.Tensor | None:\n        cache_path = self._cache_path(image_path, workspace)\n        if cache_path.exists():\n            try:\n                return torch.from_numpy(np.load(cache_path, allow_pickle=False))\n            except Exception as exc:\n                LOGGER.warning(\"Discarding invalid ColPali cache %s: %s\", cache_path.name, exc)\n                cache_path.unlink(missing_ok=True)\n        vector = self._encode_image(image_path)\n        if vector is not None:\n            np.save(cache_path, vector.numpy())\n        return vector\n\n    @torch.no_grad()\n    def _query_multivector(self, query: str) -> torch.Tensor | None:\n        try:\n            batch = self.processor.process_queries([query]).to(self.model.device)\n            return self.model(**batch).squeeze(0).to(torch.float16).cpu()\n        except Exception as exc:\n            LOGGER.error(\"ColPali query encoding failed: %s\", exc)\n            return None\n\n    def embed_image(self, image_path: str, workspace: Any = None) -> np.ndarray | None:\n        multi = (\n            self._page_multivector(image_path, workspace)\n            if workspace is not None else self._encode_image(image_path)\n        )\n        return None if multi is None else multi.float().mean(dim=0).numpy()\n\n    def embed_query(self, query: str) -> np.ndarray | None:\n        multi = self._query_multivector(query)\n        return None if multi is None else multi.float().mean(dim=0).numpy()\n\n    def _pages(self, project_id: str | None, workspace: Any) -> list[dict[str, Any]]:\n        pages: list[dict[str, Any]] = []\n        for doc_id, document in workspace.load_metadata().get(\"documents\", {}).items():\n            doc_project = document.get(\"project_id\")\n            if project_id and doc_project != project_id:\n                continue\n            for page in document.get(\"pages\", []):\n                page_num = int(page.get(\"page_num\", 1))\n                pages.append({\n                    \"id\": f\"{doc_id}_p{page_num}\",\n                    \"doc_name\": document.get(\"name\", doc_id),\n                    \"doc_id\": doc_id,\n                    \"page\": page_num,\n                    \"image_path\": workspace.resolve_page_path(page.get(\"image_path\", \"\")),\n                    \"text_content\": page.get(\"text_content\", \"\"),\n                    \"project_id\": doc_project,\n                    \"source_type\": document.get(\"source_type\", \"project_document\"),\n                })\n        return pages\n\n    def retrieve(self, query: str, top_k: int = 5, project_id: str | None = None,\n                 workspace: Any = None) -> list[dict[str, Any]]:\n        if workspace is None:\n            raise RuntimeError(\"Hybrid retrieval requires the caller's account workspace\")\n        collection = workspace.collection\n        pages = self._pages(project_id, workspace)\n        if not pages:\n            LOGGER.warning(\"No indexed pages match project_id=%r\", project_id)\n            return []\n        by_id = {page[\"id\"]: page for page in pages}\n        candidate_limit = min(len(pages), max(12, top_k * 3))\n        lexical_raw = {page[\"id\"]: lexical_score(query, page) for page in pages}\n        lexical_ids = sorted(lexical_raw, key=lexical_raw.get, reverse=True)[:candidate_limit]\n\n        query_multi = self._query_multivector(query)\n        pooled_raw: dict[str, float] = {}\n        if query_multi is not None and collection.count() > 0:\n            pooled_query = query_multi.float().mean(dim=0).numpy()\n            try:\n                kwargs: dict[str, Any] = {\n                    \"query_embeddings\": [pooled_query.tolist()],\n                    \"n_results\": min(candidate_limit, collection.count()),\n                    \"include\": [\"metadatas\", \"distances\"],\n                }\n                if project_id:\n                    kwargs[\"where\"] = {\"project_id\": project_id}\n                result = collection.query(**kwargs)\n                for item_id, distance in zip(result[\"ids\"][0], result[\"distances\"][0]):\n                    if item_id in by_id:\n                        pooled_raw[item_id] = 1.0 - float(distance)\n            except Exception as exc:\n                LOGGER.warning(\"Chroma candidate retrieval failed: %s\", exc)\n\n        preliminary_ids = set(lexical_ids) | set(pooled_raw)\n        if len(preliminary_ids) > candidate_limit:\n            lexical_norm = normalize_scores({key: lexical_raw.get(key, 0.0) for key in preliminary_ids})\n            pooled_norm = normalize_scores({key: pooled_raw.get(key, 0.0) for key in preliminary_ids})\n            preliminary_ids = set(sorted(\n                preliminary_ids,\n                key=lambda key: 0.75 * lexical_norm.get(key, 0.0) + 0.25 * pooled_norm.get(key, 0.0),\n                reverse=True,\n            )[:candidate_limit])\n        candidates = [by_id[item_id] for item_id in preliminary_ids]\n\n        late_raw: dict[str, float] = {}\n        if query_multi is not None:\n            vectors, vector_ids = [], []\n            for page in candidates:\n                image_path = page.get(\"image_path\", \"\")\n                if image_path and os.path.exists(image_path):\n                    vector = self._page_multivector(image_path, workspace)\n                    if vector is not None:\n                        vectors.append(vector)\n                        vector_ids.append(page[\"id\"])\n            if vectors:\n                try:\n                    scores = self.processor.score_multi_vector(\n                        [query_multi], vectors, batch_size=8, device=self.device\n                    )[0].tolist()\n                    late_raw = {item_id: float(score) for item_id, score in zip(vector_ids, scores)}\n                except Exception as exc:\n                    LOGGER.warning(\"ColPali MaxSim reranking failed: %s\", exc)\n\n        lexical = normalize_scores({key: lexical_raw.get(key, 0.0) for key in preliminary_ids})\n        pooled = normalize_scores({key: pooled_raw.get(key, 0.0) for key in preliminary_ids})\n        late = normalize_scores({key: late_raw.get(key, 0.0) for key in preliminary_ids})\n        strong_exact_match = max(lexical_raw.values(), default=0.0) >= 6.0\n        weights = {\n            \"lexical\": (0.70 if strong_exact_match else 0.55) if lexical_raw else 0.0,\n            \"late\": (0.28 if strong_exact_match else 0.40) if late_raw else 0.0,\n            \"pooled\": (0.02 if strong_exact_match else 0.05) if pooled_raw else 0.0,\n        }\n        total_weight = sum(weights.values()) or 1.0\n        ranked = []\n        for page in candidates:\n            item_id = page[\"id\"]\n            score = (\n                weights[\"lexical\"] * lexical.get(item_id, 0.0)\n                + weights[\"late\"] * late.get(item_id, 0.0)\n                + weights[\"pooled\"] * pooled.get(item_id, 0.0)\n            ) / total_weight\n            lexical_value = lexical.get(item_id, 0.0)\n            late_value = late.get(item_id, 0.0)\n            pooled_value = pooled.get(item_id, 0.0)\n            methods = []\n            if lexical_raw.get(item_id, 0.0) > 0:\n                methods.append(\"text match\")\n            if item_id in late_raw:\n                methods.append(\"ColPali MaxSim\")\n            if item_id in pooled_raw:\n                methods.append(\"vector candidate\")\n            haystack = f\"{page.get('doc_name', '')} {page.get('text_content', '')}\".casefold()\n            ranked.append({\n                **page,\n                \"score\": round(max(0.0, min(1.0, score)), 4),\n                \"retrieval_method\": \" + \".join(methods) or \"metadata fallback\",\n                \"score_breakdown\": {\n                    \"text\": round(lexical_value, 4),\n                    \"colpali\": round(late_value, 4),\n                    \"vector\": round(pooled_value, 4),\n                },\n                \"matched_terms\": [\n                    token for token in dict.fromkeys(search_tokens(query)) if token in haystack\n                ],\n            })\n        ranked.sort(key=lambda page: page[\"score\"], reverse=True)\n        selected = ranked[: min(top_k, len(ranked))]\n        if selected:\n            LOGGER.info(\n                \"Hybrid retrieval selected %d pages; top=%s p%d score=%.3f\",\n                len(selected), selected[0][\"doc_name\"], selected[0][\"page\"], selected[0][\"score\"],\n            )\n        torch.cuda.empty_cache()\n        return selected\n\n\ndef install_hybrid_retrieval(namespace: dict[str, Any]) -> HybridColPaliRetriever:\n    \"\"\"Replace the pooled-vector helpers with the hybrid retriever.\n\n    ``retrieve_context``, ``embed_image``, and ``embed_query`` are looked up as\n    globals at call time, so rebinding them here upgrades every existing caller.\n    Both signatures keep the ``workspace`` keyword the callers already pass.\n    \"\"\"\n    retriever = HybridColPaliRetriever(namespace)\n    namespace[\"HYBRID_RETRIEVER\"] = retriever\n    namespace[\"embed_image\"] = retriever.embed_image\n    namespace[\"embed_query\"] = retriever.embed_query\n    namespace[\"retrieve_context\"] = retriever.retrieve\n    return retriever\n", encoding="utf-8"
    )
    (integration_dir / "evidence_viewer.py").write_text(
        "\"\"\"Exact-page evidence inspection and relevance feedback for BuildMarshalAI.\"\"\"\n\nfrom __future__ import annotations\n\nimport json\nimport re\nimport uuid\nfrom datetime import datetime, timezone\nfrom pathlib import Path\nfrom typing import Any, Mapping\n\nfrom fastapi import Depends, HTTPException\nfrom pydantic import BaseModel, Field\n\n\nEVIDENCE_STOPWORDS = {\n    \"a\", \"about\", \"an\", \"and\", \"are\", \"for\", \"from\", \"give\", \"i\", \"in\",\n    \"is\", \"me\", \"of\", \"on\", \"please\", \"show\", \"tell\", \"the\", \"this\", \"to\",\n    \"very\", \"what\", \"with\", \"you\",\n}\n\n\ndef evidence_terms(query: str) -> list[str]:\n    \"\"\"Return stable, useful query terms for evidence highlighting.\"\"\"\n    return list(dict.fromkeys(\n        token for token in re.findall(r\"[a-z0-9]+\", str(query).casefold())\n        if token not in EVIDENCE_STOPWORDS\n    ))\n\n\ndef best_evidence_snippet(text: str, query: str, max_chars: int = 1400) -> str:\n    \"\"\"Select the densest query-matching window from a page's extracted text.\"\"\"\n    clean = re.sub(r\"\\s+\", \" \", str(text or \"\")).strip()\n    if len(clean) <= max_chars:\n        return clean\n    terms = evidence_terms(query)\n    if not terms:\n        return clean[:max_chars].rstrip() + \"\u2026\"\n\n    lowered = clean.casefold()\n    candidates: list[int] = []\n    for term in terms:\n        candidates.extend(match.start() for match in re.finditer(rf\"\\b{re.escape(term)}\\b\", lowered))\n    if not candidates:\n        return clean[:max_chars].rstrip() + \"\u2026\"\n\n    best_start, best_score = 0, -1\n    half = max_chars // 2\n    for position in candidates:\n        start = max(0, min(position - half, len(clean) - max_chars))\n        window = lowered[start:start + max_chars]\n        score = sum(len(re.findall(rf\"\\b{re.escape(term)}\\b\", window)) for term in terms)\n        if score > best_score:\n            best_start, best_score = start, score\n\n    if best_start:\n        boundary = clean.find(\" \", best_start)\n        if 0 <= boundary < best_start + 80:\n            best_start = boundary + 1\n    snippet = clean[best_start:best_start + max_chars].strip()\n    return (\"\u2026\" if best_start else \"\") + snippet + (\"\u2026\" if best_start + max_chars < len(clean) else \"\")\n\n\nclass EvidenceFeedbackRequest(BaseModel):\n    doc_id: str = Field(min_length=1, max_length=200)\n    page: int = Field(ge=1)\n    rating: str = Field(pattern=r\"^(relevant|not_relevant)$\")\n    query: str = Field(default=\"\", max_length=4000)\n    message_id: str | None = Field(default=None, max_length=200)\n    score: float | None = Field(default=None, ge=0.0, le=1.0)\n\n\ndef register_evidence_viewer_routes(namespace: Mapping[str, Any]) -> dict[str, Any]:\n    \"\"\"Register evidence-detail and relevance-feedback routes on the notebook app.\n\n    Both the document lookup and the feedback log resolve through the caller's\n    account workspace, so a document id from another account reads as missing\n    rather than as someone else's page.\n    \"\"\"\n    required = (\"app\", \"require_account\")\n    missing = [name for name in required if name not in namespace]\n    if missing:\n        raise RuntimeError(f\"Evidence viewer integration is missing: {', '.join(missing)}\")\n\n    app = namespace[\"app\"]\n    require_account = namespace[\"require_account\"]\n\n    def find_page(workspace: Any, doc_id: str, page_num: int) -> tuple[Mapping[str, Any], Mapping[str, Any]]:\n        document = workspace.load_metadata().get(\"documents\", {}).get(doc_id)\n        if not document:\n            raise HTTPException(status_code=404, detail=\"Document not found\")\n        for page in document.get(\"pages\", []):\n            if int(page.get(\"page_num\", 0)) == page_num:\n                return document, page\n        raise HTTPException(status_code=404, detail=\"Document page not found\")\n\n    def read_feedback(workspace: Any) -> list[dict[str, Any]]:\n        path = workspace.evidence_file\n        if not path.exists():\n            return []\n        try:\n            loaded = json.loads(path.read_text(encoding=\"utf-8\"))\n        except (OSError, json.JSONDecodeError):\n            return []\n        return loaded if isinstance(loaded, list) else []\n\n    @app.get(\"/api/evidence/{doc_id}/{page_num}\")\n    async def evidence_detail(\n        doc_id: str, page_num: int, query: str = \"\", context=Depends(require_account),\n    ) -> dict[str, Any]:\n        workspace = context.workspace\n        document, page = find_page(workspace, doc_id, page_num)\n        text = str(page.get(\"text_content\", \"\"))\n        return {\n            \"doc_id\": doc_id,\n            \"doc_name\": document.get(\"name\", doc_id),\n            \"page\": page_num,\n            \"page_count\": int(document.get(\"page_count\", len(document.get(\"pages\", [])))),\n            \"project_id\": document.get(\"project_id\"),\n            \"source_type\": document.get(\"source_type\", \"project_document\"),\n            \"text_content\": text,\n            \"evidence_text\": best_evidence_snippet(text, query),\n            \"matched_terms\": evidence_terms(query),\n            \"has_page_image\": bool(workspace.resolve_page_path(page.get(\"image_path\"))),\n            \"image_endpoint\": f\"/api/pages/{doc_id}/{page_num}\",\n        }\n\n    @app.post(\"/api/evidence/feedback\")\n    async def evidence_feedback(\n        body: EvidenceFeedbackRequest, context=Depends(require_account),\n    ) -> dict[str, Any]:\n        workspace = context.workspace\n        document, _ = find_page(workspace, body.doc_id, body.page)\n        records = read_feedback(workspace)\n        record = {\n            \"id\": uuid.uuid4().hex,\n            \"doc_id\": body.doc_id,\n            \"doc_name\": document.get(\"name\", body.doc_id),\n            \"page\": body.page,\n            \"rating\": body.rating,\n            \"query\": body.query,\n            \"message_id\": body.message_id,\n            \"score\": body.score,\n            \"user_id\": context.user_id,\n            \"created_at\": datetime.now(timezone.utc).isoformat(),\n        }\n        records.append(record)\n        records = records[-5000:]\n        path = workspace.evidence_file\n        temporary = path.with_name(path.name + \".tmp\")\n        temporary.write_text(json.dumps(records, indent=2), encoding=\"utf-8\")\n        temporary.replace(path)\n        return {\"status\": \"saved\", \"feedback_id\": record[\"id\"]}\n\n    @app.get(\"/api/evidence-feedback/stats\")\n    async def evidence_feedback_stats(context=Depends(require_account)) -> dict[str, Any]:\n        records = read_feedback(context.workspace)\n        relevant = sum(item.get(\"rating\") == \"relevant\" for item in records)\n        not_relevant = sum(item.get(\"rating\") == \"not_relevant\" for item in records)\n        return {\"total\": len(records), \"relevant\": relevant, \"not_relevant\": not_relevant}\n\n    return {\"scoped\": \"per-account\"}\n\n\n__all__ = [\n    \"EvidenceFeedbackRequest\", \"best_evidence_snippet\", \"evidence_terms\",\n    \"register_evidence_viewer_routes\",\n]\n", encoding="utf-8"
    )
    repository = "https://raw.githubusercontent.com/shafitanvir32/BuildMarshalAI_capstone/docgen-pipeline/backend"
    for filename in ("accounts.py", "external_imports.py", "oauth_tokens.py",
                     "tasks.py", "project_management.py", "company_settings.py",
                     "permissions.py", "user_roles.py",
                     "todo_lists.py",
                     "entity_schema.py",
                     "project_analytics.py",
                     "project_report.py",
                     "document_storage.py",
                     "project_onboarding.py",
                     "chat_entities.py",
                     "meeting_language.py",
                     "document_generation.py", "google_workspace.py",
                     "microsoft_workspace.py"):
        urllib.request.urlretrieve(f"{repository}/{filename}", integration_dir / filename)
else:
    candidates = [
        Path.cwd() / "backend",
        Path.cwd(),
        Path.cwd() / "BuildMarshalAI_capstone" / "backend",
        Path.cwd() / "BuildMarshalAI_capstone",
    ]
    integration_dir = next(
        (candidate for candidate in candidates
         if (candidate / "accounts.py").exists()
         and (candidate / "tasks.py").exists()
         and (candidate / "project_management.py").exists()
         and (candidate / "company_settings.py").exists()
         and (candidate / "permissions.py").exists()
         and (candidate / "user_roles.py").exists()
         and (candidate / "todo_lists.py").exists()
         and (candidate / "entity_schema.py").exists()
         and (candidate / "project_analytics.py").exists()
         and (candidate / "project_report.py").exists()
         and (candidate / "document_storage.py").exists()
         and (candidate / "project_onboarding.py").exists()
         and (candidate / "chat_entities.py").exists()
         and (candidate / "meeting_language.py").exists()
         and (candidate / "oauth_tokens.py").exists()
         and (candidate / "external_imports.py").exists()
         and (candidate / "document_generation.py").exists()
         and (candidate / "google_workspace.py").exists()
         and (candidate / "microsoft_workspace.py").exists()
         and (candidate / "hybrid_retrieval.py").exists()
         and (candidate / "evidence_viewer.py").exists()),
        None,
    )
    if integration_dir is None:
        raise RuntimeError(
            "Run this notebook from the BuildMarshalAI repository root or backend directory."
        )

sys.path.insert(0, str(integration_dir))
from accounts import register_account_routes
from tasks import register_task_routes
from project_management import register_project_management_routes
from company_settings import register_company_settings_routes
from user_roles import register_user_role_routes
from todo_lists import register_todo_list_routes
from project_onboarding import register_project_onboarding_routes
from chat_entities import register_chat_entity_routes
from project_analytics import register_project_analytics_routes
from project_report import register_project_report_routes
from document_storage import register_document_storage_routes
from hybrid_retrieval import install_hybrid_retrieval
from evidence_viewer import register_evidence_viewer_routes
from document_generation import register_kaggle_routes
from google_workspace import register_google_workspace_routes
from microsoft_workspace import register_microsoft_workspace_routes

# Accounts register first: the remaining modules resolve their per-request
# workspace through the context this call publishes, and any legacy
# single-tenant data is migrated into an owned account here.
ACCOUNT_SERVICE = register_account_routes(globals())
TASK_SERVICE = register_task_routes(globals())
PROJECT_MANAGEMENT_SERVICE = register_project_management_routes(globals())
USER_ROLE_SERVICE = register_user_role_routes(globals())
COMPANY_SETTINGS_SERVICE = register_company_settings_routes(globals())
HYBRID_RETRIEVER = install_hybrid_retrieval(globals())
EVIDENCE_VIEWER_SERVICE = register_evidence_viewer_routes(globals())
DOCUMENT_GENERATION_SERVICE = register_kaggle_routes(globals())
GOOGLE_WORKSPACE_SERVICE = register_google_workspace_routes(globals())
MICROSOFT_WORKSPACE_SERVICE = register_microsoft_workspace_routes(globals())
TODO_LIST_SERVICE = register_todo_list_routes(globals())
# Onboarding registers last: it reuses the project, task, catalogue, and role
# constructors the modules above install, and writes nothing live until commit.
PROJECT_ONBOARDING_SERVICE = register_project_onboarding_routes(globals())
# Marshal Chat can create a project or a user from a sentence; the schema
# decides what is mandatory and the ordinary create routes do the creating.
CHAT_ENTITY_SERVICE = register_chat_entity_routes(globals())
# Statistics, the report that reads them, and the storage the documents sit in.
PROJECT_ANALYTICS_SERVICE = register_project_analytics_routes(globals())
PROJECT_REPORT_SERVICE = register_project_report_routes(globals())
DOCUMENT_STORAGE_SERVICE = register_document_storage_routes(globals())
print(f"Account isolation active - {ACCOUNT_SERVICE['accounts']} account(s) registered")
print("Project task routes registered")
print("Project people, cost, timeline, and procurement routes registered")
print("Hybrid text + ColPali MaxSim retrieval registered")
print("Exact-page evidence viewer routes registered")
print("Project document-generation routes registered")
print("Google Workspace routes registered", "(configured)" if GOOGLE_WORKSPACE_SERVICE["configured"] else "(configure OAuth variables to enable)")
print("Microsoft 365 routes registered", "(configured)" if MICROSOFT_WORKSPACE_SERVICE["configured"] else "(configure OAuth variables to enable)")
print(f"Project onboarding routes registered - draft kinds: {', '.join(PROJECT_ONBOARDING_SERVICE['kinds'])}")
print("Project statistics routes registered -", ", ".join(PROJECT_ANALYTICS_SERVICE["phases"]))
print("Project report routes registered -", len(PROJECT_REPORT_SERVICE["sections"]), "sections")
print("Document storage routes registered -", ", ".join(DOCUMENT_STORAGE_SERVICE["actions"]))


# ======================================================================
# CELL 13
# ======================================================================

import socket
import uvicorn
from pyngrok import ngrok


def resolve_backend_port():
    """The port to serve on, which is not always the one that was asked for.

    Windows hands blocks of low TCP ports to Hyper-V at every boot and then
    refuses to bind them even though nothing is listening on them, so a port
    that worked yesterday can fail today with "an attempt was made to access
    a socket in a way forbidden by its access permissions". Serving on a port
    that works beats refusing to start, and the browser is told where to look
    by publish_backend_port below.
    """
    wanted = int(os.environ.get("BUILDMARSHAL_PORT", "8000"))
    first_refusal = None
    for candidate in (wanted, 8900, 8901, 9000, 9100, 18000):
        # Both addresses: uvicorn binds the wildcard, and on Windows that
        # succeeds even when another process already holds 127.0.0.1 on the
        # same port -- the more specific bind then wins every loopback
        # request and the backend is up but unreachable.
        blocked = None
        for address in ("0.0.0.0", "127.0.0.1"):
            probe = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
            try:
                probe.bind((address, candidate))
            except OSError as error:
                blocked = error
            finally:
                probe.close()
            if blocked:
                break
        if blocked:
            print(f"Port {candidate} cannot be bound ({blocked.strerror}).")
            first_refusal = first_refusal or blocked
            continue
        if candidate != wanted:
            # WSAEACCES is Windows refusing a port it has reserved for
            # Hyper-V; anything else is an ordinary clash with a process.
            reserved = getattr(first_refusal, "winerror", None) == 10013
            reason = "is reserved by Windows" if reserved else "is already in use"
            print(f"Port {wanted} {reason}; serving on {candidate}.")
            if reserved:
                print("Run RESERVE-PORTS.ps1 as Administrator to get the usual port back.")
        return candidate
    raise RuntimeError("No bindable port for the backend. Run RESERVE-PORTS.ps1 "
                       "as Administrator, or reboot.")


def publish_backend_port(port):
    """Write the port where the frontend reads it: a browser has no environment."""
    here = Path(globals()["__file__"]).resolve().parent if "__file__" in globals() else Path.cwd()
    for parent in (here, *here.parents[:3]):
        frontend = parent / "frontend"
        if frontend.is_dir():
            (frontend / "local-config.js").write_text(
                "// Written by the backend at startup -- do not edit.\n"
                "// A URL entered on the sign-in screen still wins over this.\n"
                f"window.BMARSHAL_API_URL = 'http://127.0.0.1:{port}';\n",
                encoding="utf-8")
            return


def start_server():
    port = resolve_backend_port()
    publish_backend_port(port)
    public_url = None
    if NGROK_AUTH_TOKEN:
        try:
            ngrok.set_auth_token(NGROK_AUTH_TOKEN)
            for tunnel in ngrok.get_tunnels():
                ngrok.disconnect(tunnel.public_url)
            ngrok.kill()
            public_url = ngrok.connect(port).public_url
        except Exception as exc:
            print(f"ngrok could not start; continuing locally: {exc}")

    local_url = f"http://127.0.0.1:{port}"
    shown_url = public_url or local_url
    print("\n" + "=" * 60)
    print("BuildMarshalAI CLIProxyAPI Backend is LIVE")
    print("=" * 60)
    print(f"Backend URL : {shown_url}")
    print(f"API Docs    : {shown_url}/docs")
    print(f"Health      : {shown_url}/api/health")
    print(f"Proxy check : {shown_url}/api/cliproxy/status")
    print(f"Retrieval   : ColPali on {COLPALI_DEVICE}")
    print(f"Generation  : {CLIPROXY_CLIENT.base_url}")
    print("=" * 60 + "\n")
    uvicorn.run(app, host="0.0.0.0", port=port)

start_server()
